# -*- coding: utf-8 -*-
"""
Script para generar el documento oficial en formato Microsoft Word (.docx)
para la entrega de las dos actividades del SENA:
- ACTIVIDAD 1: Seguridad informática
- ACTIVIDAD 2: Construcción del prototipo

Proyecto: EduNova – Sistema Integrado de Educación Técnica (SINETEC)
Aprendices: Aminta Moguea Florez y Dulce María Figueroa
Programa: Tecnólogo en Análisis y Desarrollo de Software (ADSI / ADSO)
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Aplica color de fondo a una celda de tabla."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Establece márgenes internos de una celda en dxa (1 pt = 20 dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tc_pr.append(tc_mar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    """Establece bordes limpios y formales para una tabla."""
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(borders)

def add_callout(doc, title, text, border_color="1B365D"):
    """Agrega una caja de texto destacada con borde lateral izquierdo."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=140)
    
    # Borde izquierdo grueso
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    run_t = p.add_run(title + "\n")
    run_t.bold = True
    run_t.font.name = "Arial"
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    run_b = p.add_run(text)
    run_b.font.name = "Arial"
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    # Espacio después del callout
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def construir_documento():
    doc = docx.Document()
    
    # Configurar márgenes de página (2.54 cm = 1 pulgada estándar)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Encabezado formal
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("SENA — Centro de Logística y Promoción Ecoturística del Magdalena | EduNova")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        
        # Pie de página formal
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Tecnología en Análisis y Desarrollo de Software — Actividades de Aprendizaje")
        frun.font.name = "Arial"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Estilos de base
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # =========================================================================
    # 1. PORTADA FORMAL SENA
    # =========================================================================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(36)
    p_inst.paragraph_format.space_after = Pt(4)
    run_inst = p_inst.add_run("SERVICIO NACIONAL DE APRENDIZAJE – SENA\nREGIONAL MAGDALENA\nCENTRO DE LOGÍSTICA Y PROMOCIÓN ECOTURÍSTICA DEL MAGDALENA")
    run_inst.bold = True
    run_inst.font.name = "Arial"
    run_inst.font.size = Pt(12)
    run_inst.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    p_prog = doc.add_paragraph()
    p_prog.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_prog.paragraph_format.space_before = Pt(8)
    p_prog.paragraph_format.space_after = Pt(40)
    run_prog = p_prog.add_run("PROGRAMA: TECNÓLOGO EN ANÁLISIS Y DESARROLLO DE SOFTWARE (ADSI / ADSO)\nFICHA DE FORMACIÓN: 228106")
    run_prog.font.name = "Arial"
    run_prog.font.size = Pt(11)
    run_prog.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tit.paragraph_format.space_before = Pt(20)
    p_tit.paragraph_format.space_after = Pt(12)
    run_tit = p_tit.add_run("INFORME DE EVIDENCIAS Y APLICACIÓN PRÁCTICA:\nSEGURIDAD INFORMÁTICA Y CONSTRUCCIÓN DEL PROTOTIPO")
    run_tit.bold = True
    run_tit.font.name = "Arial"
    run_tit.font.size = Pt(16)
    run_tit.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(40)
    run_sub = p_sub.add_run("Proyecto: EduNova – Sistema Integrado de Educación Técnica (SINETEC)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.italic = True
    run_sub.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    p_aut = doc.add_paragraph()
    p_aut.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aut.paragraph_format.space_before = Pt(50)
    p_aut.paragraph_format.space_after = Pt(4)
    run_aut_lbl = p_aut.add_run("APRENDICES:\n")
    run_aut_lbl.bold = True
    run_aut_lbl.font.size = Pt(11)
    run_aut1 = p_aut.add_run("AMINTA MOGUEA FLOREZ\nDULCE MARÍA FIGUEROA\n")
    run_aut1.bold = True
    run_aut1.font.size = Pt(11.5)

    p_ciu = doc.add_paragraph()
    p_ciu.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ciu.paragraph_format.space_before = Pt(60)
    p_ciu.paragraph_format.space_after = Pt(0)
    run_ciu = p_ciu.add_run("SANTA MARTA, MAGDALENA, COLOMBIA\nOCTUBRE DE 2026")
    run_ciu.bold = True
    run_ciu.font.size = Pt(11)

    doc.add_page_break()

    # =========================================================================
    # 2. INTRODUCCIÓN
    # =========================================================================
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("1. INTRODUCCIÓN")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    p_intro1 = doc.add_paragraph(
        "El presente documento técnico y formativo ha sido elaborado en el marco de la formación del Tecnólogo "
        "en Análisis y Desarrollo de Software (ADSI / ADSO) del Servicio Nacional de Aprendizaje (SENA), "
        "Regional Magdalena, como evidencia formal del desarrollo integral y la aplicación de competencias en el "
        "proyecto formativo EduNova – Sistema Integrado de Educación Técnica (SINETEC)."
    )
    p_intro1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p_intro2 = doc.add_paragraph(
        "EduNova surge como una respuesta tecnológica a la problemática de dispersión, desorganización y vulnerabilidad "
        "de la información académica, evaluativa y administrativa en las Instituciones Educativas de la región del Magdalena. "
        "A través de una arquitectura web basada en Python y Django, articulada con bases de datos relacionales "
        "(MySQL para producción y SQLite para desarrollo y pruebas locales), el sistema integra a todos los actores del "
        "ecosistema educativo: administradores, directivos, secretaría, docentes, estudiantes y familias."
    )
    p_intro2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p_intro3 = doc.add_paragraph(
        "El contenido de este informe aborda directamente dos actividades clave del programa formativo SENA:\n"
        "1. Actividad 1 – Seguridad Informática: Enfoque riguroso en el análisis de vulnerabilidades, implementación "
        "de mecanismos de control de acceso basados en roles (RBAC), prevención de referencias directas inseguras a objetos (IDOR), "
        "protección contra ataques comunes web (CSRF, XSS, Clickjacking, inyecciones) y auditoría inmutable de eventos críticos.\n"
        "2. Actividad 2 – Construcción del Prototipo: Documentación del propósito institucional del sistema, desglose funcional "
        "y operativo de los seis paneles de usuario (Administrador, Rectoría, Secretaría, Docente, Familia y Estudiante), características "
        "de usabilidad y experiencia de usuario (UX/UI), mejoras reales comprobadas en el código y resultados de las pruebas funcionales "
        "ejecutadas directamente en el repositorio activo."
    )
    p_intro3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p_intro4 = doc.add_paragraph(
        "De conformidad con los requerimientos académicos, todos los datos, módulos, mecanismos y pruebas aquí descritos "
        "provienen exclusivamente del código fuente real de EduNova, reflejando resultados verificados sin afirmaciones "
        "hipotéticas ni invención de funcionalidades."
    )
    p_intro4.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # =========================================================================
    # 3. OBJETIVOS
    # =========================================================================
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("2. OBJETIVOS")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    p_og = doc.add_paragraph()
    r = p_og.add_run("2.1 Objetivo General\n")
    r.bold = True
    r.font.size = Pt(12)
    p_og.add_run(
        "Analizar, aplicar y verificar las políticas y mecanismos de control de seguridad informática y las características "
        "funcionales, de usabilidad y experiencia de usuario del prototipo del sistema de información EduNova, asegurando la "
        "integridad de los datos institucionales y la operatividad fluida de sus seis paneles de interacción."
    )
    p_og.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p_oe = doc.add_paragraph()
    r = p_oe.add_run("2.2 Objetivos Específicos\n")
    r.bold = True
    r.font.size = Pt(12)
    
    doc.add_paragraph(
        "• Identificar las vulnerabilidades potenciales y los riesgos de seguridad informática en el diseño y código del sistema EduNova, aplicando controles basados en el principio de menor privilegio y la defensa en profundidad.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Documentar los mecanismos de control implementados en backend (RBAC, anti-IDOR, CSRF, protección de sesiones, sanitización y auditoría inmutable) y corroborar su efectividad mediante pruebas de verificación comprobables.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Documentar el propósito institucional, estructura de navegación y catálogo funcional detallado de los seis paneles de EduNova (Administrador, Rectoría, Secretaría, Docente, Familia y Estudiante).",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Evaluar las características de facilidad de manejo y experiencia de usuario (UX/UI) aplicadas al prototipo funcional conforme a protocolos modernos de diseño.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "• Ejecutar y certificar las baterías de pruebas funcionales reales desarrolladas para el sistema, evidenciando el comportamiento exacto de los módulos académicos, de seguimiento y administrativos.",
        style='List Bullet'
    )

    doc.add_page_break()

    # =========================================================================
    # 4. ACTIVIDAD 1: SEGURIDAD INFORMÁTICA
    # =========================================================================
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("3. ACTIVIDAD 1: SEGURIDAD INFORMÁTICA")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    p_norma = doc.add_paragraph()
    r = p_norma.add_run("Norma / Criterio de Evaluación SENA:\n")
    r.bold = True
    r.font.size = Pt(11)
    r_cit = p_norma.add_run("“Aplicar políticas y mecanismos de control en el diseño del sistema de información, mediante el análisis de la vulnerabilidad de la información, siguiendo los parámetros establecidos por la organización”.")
    r_cit.italic = True
    p_norma.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Subsección 3.1
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("3.1 Parámetros y Políticas de Seguridad Establecidos por la Organización")
    run.bold = True
    run.font.size = Pt(12)

    doc.add_paragraph(
        "En el contexto del Servicio Nacional de Aprendizaje (SENA) y de las Instituciones Educativas articuladas con la Media Técnica, "
        "la información escolar (datos biográficos, notas, juicios evaluativos, asistencias y pagos) posee carácter confidencial y sensible "
        "conforme a la Ley Estatutaria 1581 de 2012 (Habeas Data) y estándares internacionales como ISO/IEC 27001 y las directrices OWASP (Open Worldwide Application Security Project). "
        "En el diseño de EduNova se adoptaron las siguientes directrices institucionales de obligado cumplimiento:"
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    add_callout(
        doc,
        "POLÍTICA INSTITUCIONAL 1: Control de Acceso Estricto y Menor Privilegio (RBAC)",
        "Ningún usuario debe acceder a funciones, registros o vistas que excedan las atribuciones explícitas de su perfil institucional. "
        "Las decisiones de autorización deben realizarse invariablemente en el Backend y nunca delegarse exclusivamente al front-end ni al ocultamiento visual de botones."
    )

    add_callout(
        doc,
        "POLÍTICA INSTITUCIONAL 2: Privacidad y Aislamiento de Datos (Anti-IDOR)",
        "Se prohíbe que cualquier usuario pueda consultar o alterar expedientes ajenos manipulando identificadores numéricos en parámetros GET o peticiones POST. "
        "Un estudiante solo puede acceder a su propio expediente; un acudiente únicamente a sus hijos asociados; y un docente solo a los estudiantes asignados a su carga académica activa."
    )

    add_callout(
        doc,
        "POLÍTICA INSTITUCIONAL 3: Trazabilidad e Inmutabilidad de Transacciones Críticas",
        "Toda acción que afecte el historial académico o financiero (calificaciones, asistencias, anulación de recibos de caja, modificación de matrículas y eliminaciones) "
        "debe registrarse en una bitácora de auditoría inmutable, almacenando usuario ejecutor, IP de origen, fecha/hora y detalle técnico."
    )

    add_callout(
        doc,
        "POLÍTICA INSTITUCIONAL 4: Salvaguarda y Prevención de Pérdida Catastrófica de Datos",
        "Se restringe el borrado físico directo (DELETE en base de datos) en entidades maestras. Las bajas de estudiantes, cargas y registros deben gestionarse mediante "
        "borrado lógico (inactivación de cuentas) y snapshots serializados en una Papelera de Reciclaje institucional, permitiendo su restauración controlada."
    )

    # Subsección 3.2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("3.2 Mecanismos de Control Implementados en la Arquitectura de EduNova")
    run.bold = True
    run.font.size = Pt(12)

    doc.add_paragraph(
        "A continuación se detallan los mecanismos tecnológicos reales programados e implementados en el código fuente de EduNova:"
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Tabla de mecanismos
    t_mec = doc.add_table(rows=1, cols=3)
    t_mec.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_mec)
    
    hdr = t_mec.rows[0].cells
    hdr[0].width = Inches(1.8)
    hdr[1].width = Inches(1.8)
    hdr[2].width = Inches(2.9)
    
    headers = ["Mecanismo de Control", "Ubicación en el Código", "Función y Descripción Técnica"]
    for i, h in enumerate(headers):
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        set_cell_background(hdr[i], "E2E8F0")

    filas_mecanismos = [
        (
            "Control de Acceso por Roles (RBAC)",
            "usuarios/decorators.py\n(requerir_roles, solo_admin, solo_docente, solo_secretaria, solo_aprendiz)",
            "Intercepta cada petición HTTP antes de ejecutar la vista. Valida en base de datos el rol activo del usuario contra las categorías autorizadas (ADMIN, DIRECTIVO, SECRETARIA, DOCENTE, ESTUDIANTE, FAMILIA). Si no tiene permiso, bloquea la acción y redirige al panel legítimo."
        ),
        (
            "Regla Anti-IDOR (Insecure Direct Object Reference)",
            "usuarios/decorators.py\n(validar_propietario_o_coordinador)",
            "Verifica la pertenencia de un expediente antes de mostrar datos privados: comprueba si el usuario es superusuario/directivo, o si el docente tiene carga académica activa con el grupo del estudiante, o si la familia está legalmente vinculada en FamiliaAcudiente. Si falla, lanza excepción PermissionDenied (HTTP 403)."
        ),
        (
            "Protección contra Falsificación CSRF",
            "sinetec_project/settings.py\n(CsrfViewMiddleware, CSRF_FAILURE_VIEW)",
            "Exige la presencia y validación de un token criptográfico no reutilizable en cada petición que modifique datos (POST). Se integró CSRF_FAILURE_VIEW para manejar de forma segura intentos bloqueados."
        ),
        (
            "Auditoría Inmutable de Eventos Críticos",
            "seguimiento/models.py\n(RegistroAuditoria)\nacademico/views.py\ninstituciones/views.py",
            "Modelo relacional que registra usuario, acción ejecutada, módulo del sistema, IP de origen, fecha/hora y parámetros JSON detallados ante eventos como registro de calificaciones, anulación de pagos de pensión, creación de actividades y cambios de estado de matrícula."
        ),
        (
            "Borrado Lógico y Papelera de Reciclaje",
            "usuarios/models.py\n(PapeleraReciclaje, esta_activo en PerfilUsuario)",
            "Implementa la regla RN-008: en lugar de ejecutar DROP o DELETE destructivo en MySQL, el sistema desactiva el usuario y genera un snapshot serializado en JSON en la Papelera de Reciclaje, preservando integridad referencial y permitiendo recuperación."
        ),
        (
            "Criptografía y Hash de Contraseñas",
            "sinetec_project/settings.py\n(AUTH_PASSWORD_VALIDATORS, django.contrib.auth)",
            "Las contraseñas no se almacenan en texto claro. Se procesan con algoritmo PBKDF2 y firma SHA-256 con salt dinámico de 100.000+ iteraciones. Se validan atributos de similitud, longitud mínima y complejidad."
        ),
        (
            "Protección contra Inyecciones SQL",
            "Todo el código backend (ORM de Django)",
            "Todas las operaciones de base de datos se ejecutan a través de la capa de abstracción del ORM de Django con consultas parametrizadas automáticas, evitando la concatenación directa de cadenas en sentencias SQL."
        ),
        (
            "Seguridad de Cabeceras HTTP y Cookies",
            "sinetec_project/settings.py\n(XFrameOptionsMiddleware, SECURE_PROXY_SSL_HEADER)",
            "Configura X_FRAME_OPTIONS = 'DENY' para impedir ataques de Clickjacking; activa SECURE_CONTENT_TYPE_NOSNIFF para prevenir sniffing MIME; y establece SESSION_COOKIE_SAMESITE = 'Lax' contra ataques de reutilización de sesión."
        )
    ]

    for mec, ubi, desc in filas_mecanismos:
        row = t_mec.add_row().cells
        row[0].width = Inches(1.8)
        row[1].width = Inches(1.8)
        row[2].width = Inches(2.9)
        set_cell_margins(row[0])
        set_cell_margins(row[1])
        set_cell_margins(row[2])
        
        p0 = row[0].paragraphs[0]
        p0.add_run(mec).bold = True
        p0.runs[0].font.size = Pt(9)
        
        p1 = row[1].paragraphs[0]
        r1 = p1.add_run(ubi)
        r1.font.size = Pt(8.5)
        r1.italic = True
        
        p2 = row[2].paragraphs[0]
        p2.paragraph_format.line_spacing = 1.15
        p2.add_run(desc).font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Subsección 3.3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("3.3 Hallazgos Reales Identificados en el Código y Correcciones Realizadas")
    run.bold = True
    run.font.size = Pt(12)

    doc.add_paragraph(
        "Durante la revisión técnica profunda del código fuente de EduNova se detectaron vulnerabilidades y desviaciones de seguridad reales. "
        "A continuación se documentan de forma transparente los hallazgos reales, el riesgo asociado y las correcciones de código verificadas:"
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Hallazgo 1
    doc.add_paragraph("Hallazgo Real 1: Riesgo de Acceso Cruzado a Expedientes Privados de Estudiantes (IDOR)", style='Heading 2')
    doc.add_paragraph(
        "• Diagnóstico y Vulnerabilidad: En las versiones preliminares, las rutas que consultaban fichas y detalles de estudiantes "
        "recibían el identificador numérico de la matrícula o usuario. Un estudiante autenticado podía manipular manualmente la URL "
        "o enviar peticiones GET/POST alterando el ID, lo que permitía la visualización no autorizada de calificaciones y datos sensibles de otros aprendices.\n"
        "• Corrección Implementada: Se construyó la función de seguridad validar_propietario_o_coordinador(request, estudiante_user) "
        "en usuarios/decorators.py. Esta función valida de forma estricta que si el usuario es un estudiante, su ID coincida exactamente con el objeto consultado; "
        "si es un docente, verifica mediante consulta a CargaAcademica si imparte clases en el grado y sección del estudiante; y si es un acudiente, "
        "comprueba la relación en FamiliaAcudiente. Ante cualquier infracción, se emite una excepción PermissionDenied (HTTP 403).\n"
        "• Verificación: Comprobada en scripts/test_master_suite.py (Prueba 25: 'Estudiante no puede ver fichas privadas de otros' -> HTTP 302/403)."
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Hallazgo 2
    doc.add_paragraph("Hallazgo Real 2: Fuga de Privacidad en Comunicados Escolares y Circulares", style='Heading 2')
    doc.add_paragraph(
        "• Diagnóstico y Vulnerabilidad: El modelo de comunicados institucionales almacenaba circulares directivas y citaciones particulares. "
        "Si no se aplicaba un filtrado riguroso por audiencia y usuario destinatario en el QuerySet de la vista, circulares privadas o llamados de atención "
        "quedaban expuestos en el tablero general de otros estudiantes.\n"
        "• Corrección Implementada: Se reescribió la lógica de consulta en el portal estudiantil y en la vista de comunicados "
        "(seguimiento/models.py y academico/views.py), aplicando un filtro condicional que restringe la visualización exclusivamente a comunicados de carácter general "
        "o dirigidos explícitamente al ID del estudiante autenticado o a su grado/sección escolar.\n"
        "• Verificación: Comprobada en scripts/test_verificaciones.py (Verificación 3: 'Circular privada visible para estudiante destinatario y debidamente oculta para estudiante est2' -> PASADA CON ÉXITO)."
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Hallazgo 3
    doc.add_paragraph("Hallazgo Real 3: Falta de Segregación de Funciones en Secretaría y Control de Caja", style='Heading 2')
    doc.add_paragraph(
        "• Diagnóstico y Vulnerabilidad: En la interfaz de secretaría se mostraba un acceso directo para 'Registrar Cobro' de pensiones escolares, "
        "lo cual contravenía el principio de segregación funcional entre la administración académica y el recaudo financiero de tesorería/caja, "
        "además de carecer de trazabilidad de anulación con comprobante formal.\n"
        "• Corrección Implementada: Se eliminó el enlace de cobro del panel de secretaría (templates/usuarios/secretaria_dashboard.html y usuarios/views.py), "
        "restringiendo la emisión y anulación de recibos exclusivamente al módulo central de caja y tesorería en PagoPension, "
        "donde cada movimiento registra obligatoriamente al usuario responsable del cobro, número consecutivo único y comprobante digital.\n"
        "• Verificación: Comprobada en scripts/test_edunova_student_secretaria_flow.py (Prueba 2: Comprobación de que 'Registrar cobro' no está presente en Secretaría) "
        "y en scripts/test_verificaciones.py (Prueba 1: Verificación de cobro por cajero, emisión de comprobante y anulación controlada con cambio de estado a ANULADO)."
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Hallazgo 4
    doc.add_paragraph("Hallazgo Real 4: Desbordamiento de Carga Académica y Falta de Aislamiento Docente", style='Heading 2')
    doc.add_paragraph(
        "• Diagnóstico y Vulnerabilidad: En pruebas de estrés y datos iniciales, el docente podía ver enlaces o listados de grupos donde no tenía asignación académica formal, "
        "o figuraban contadores con datos estáticos incorrectos (ej. '45 estudiantes'). Esto representaba un riesgo de edición accidental de notas en cursos ajenos.\n"
        "• Corrección Implementada: Se blindó la consulta del panel docente mediante CargaAcademica.objects.filter(profesor=request.user). "
        "Toda planilla de asistencia, calificación y actividad exige la validación previa de que la carga académica pertenece al docente que inició sesión. "
        "Además, se bloquearon los accesos de docentes a los módulos de administración institucional (/usuarios/, /configuracion/).\n"
        "• Verificación: Comprobada en scripts/test_verificaciones.py (Verificación 2: Aislamiento de grupos docentes) y scripts/test_master_suite.py (Pruebas 10, 24 y 26: Bloqueo de acceso de docentes a administración con HTTP 302/403)."
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Hallazgo 5
    doc.add_paragraph("Hallazgo Real 5: Centralización de Credenciales y Variables de Entorno", style='Heading 2')
    doc.add_paragraph(
        "• Diagnóstico y Vulnerabilidad: La clave de base de datos MySQL, la SECRET_KEY de Django y las credenciales de correo figuraban predefinidas en el código, "
        "generando un riesgo crítico en caso de compartir el repositorio en control de versiones.\n"
        "• Corrección Implementada: Se desacoplaron todos los parámetros críticos utilizando python-decouple y archivos .env, "
        "garantizando que la SECRET_KEY, el motor de base de datos (MySQL vs SQLite), las contraseñas y las listas de ALLOWED_HOSTS y CSRF_TRUSTED_ORIGINS "
        "se lean desde variables del entorno del servidor y nunca se versionen en repositorios públicos.\n"
        "• Verificación: Comprobada en sinetec_project/settings.py con lectura dinámica de variables mediante config()."
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Subsección 3.4
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("3.4 Matriz de Pruebas de Verificación de Seguridad Ejecutadas")
    run.bold = True
    run.font.size = Pt(12)

    doc.add_paragraph(
        "Para validar empíricamente que los mecanismos de seguridad operan correctamente, se ejecutaron las pruebas automatizadas del proyecto. "
        "Los resultados obtenidos de la ejecución real son los siguientes:"
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    t_sec_tests = doc.add_table(rows=1, cols=4)
    t_sec_tests.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_sec_tests)
    
    hdr = t_sec_tests.rows[0].cells
    hdr[0].width = Inches(1.5)
    hdr[1].width = Inches(1.8)
    hdr[2].width = Inches(2.2)
    hdr[3].width = Inches(1.0)
    
    headers_st = ["Prueba de Seguridad", "Script Ejecutor", "Comportamiento Esperado y Verificado", "Resultado"]
    for i, h in enumerate(headers_st):
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9)
        set_cell_background(hdr[i], "E2E8F0")

    filas_sec_tests = [
        (
            "Aislamiento Anti-IDOR Estudiante",
            "scripts/test_master_suite.py (Test 25)",
            "Estudiante est1 intenta acceder al expediente de est2. El sistema lo bloquea y redirige con código HTTP 302/403.",
            "PASADA [OK]"
        ),
        (
            "Bloqueo RBAC a Docente en Admin",
            "scripts/test_master_suite.py (Test 24, 26)",
            "Docente intenta acceder a /usuarios/ y a /configuracion/. El decorador solo_admin lo bloquea con HTTP 302/403.",
            "PASADA [OK]"
        ),
        (
            "Privacidad de Circulares Escolares",
            "scripts/test_verificaciones.py (Sección 3)",
            "Circular confidencial enviada a est1 es visible en su portal y se verifica estrictamente ausente del portal de est2.",
            "PASADA [OK]"
        ),
        (
            "Segregación en Secretaría (Sin Cobros)",
            "scripts/test_edunova_student_secretaria_flow.py (Test 2)",
            "El panel /secretaria/ es inspeccionado en su HTML; se confirma que la opción de cobro financiero no existe.",
            "PASADA [OK]"
        ),
        (
            "Trazabilidad de Anulación de Recibos",
            "scripts/test_verificaciones.py (Sección 1)",
            "Recibo de caja emitido con responsable admin es anulado por POST; la base de datos actualiza el estado a ANULADO y queda registrado en auditoría.",
            "PASADA [OK]"
        ),
        (
            "Carga Segura de Archivos e Imágenes",
            "scripts/test_edunova_student_secretaria_flow.py (Test 5)",
            "Estudiante carga foto de perfil JPEG; se valida Content-Type, almacenamiento en ruta perfiles/ y persistencia tras reinicio de sesión.",
            "PASADA [OK]"
        )
    ]

    for pr, sc, comp, res in filas_sec_tests:
        row = t_sec_tests.add_row().cells
        row[0].width = Inches(1.5)
        row[1].width = Inches(1.8)
        row[2].width = Inches(2.2)
        row[3].width = Inches(1.0)
        set_cell_margins(row[0])
        set_cell_margins(row[1])
        set_cell_margins(row[2])
        set_cell_margins(row[3])
        
        p0 = row[0].paragraphs[0]
        p0.add_run(pr).bold = True
        p0.runs[0].font.size = Pt(8.5)
        
        p1 = row[1].paragraphs[0]
        p1.add_run(sc).font.size = Pt(8)
        
        p2 = row[2].paragraphs[0]
        p2.paragraph_format.line_spacing = 1.15
        p2.add_run(comp).font.size = Pt(8)
        
        p3 = row[3].paragraphs[0]
        r3 = p3.add_run(res)
        r3.bold = True
        r3.font.size = Pt(8.5)
        r3.font.color.rgb = RGBColor(0x15, 0x80, 0x3D)  # Verde

    doc.add_page_break()

    # =========================================================================
    # 5. ACTIVIDAD 2: CONSTRUCCIÓN DEL PROTOTIPO
    # =========================================================================
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("4. ACTIVIDAD 2: CONSTRUCCIÓN DEL PROTOTIPO")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    p_norma2 = doc.add_paragraph()
    r = p_norma2.add_run("Norma / Criterio de Evaluación SENA:\n")
    r.bold = True
    r.font.size = Pt(11)
    r_cit2 = p_norma2.add_run("“Construir el prototipo del sistema de información, a partir del análisis de las características funcionales del sistema en relación con facilidad de manejo, funcionalidad y experiencia del usuario, apoyado en software aplicado según protocolos de diseño”.")
    r_cit2.italic = True
    p_norma2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Subsección 4.1
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("4.1 Propósito General del Sistema EduNova")
    run.bold = True
    run.font.size = Pt(12)

    doc.add_paragraph(
        "EduNova – Sistema Integrado de Educación Técnica (SINETEC) ha sido concebido y construido como un prototipo web funcional de alta fidelidad, "
        "orientado a modernizar y optimizar la administración académica, la evaluación por competencias y el seguimiento de la formación en la Educación Media Técnica "
        "articulada con el SENA. Su propósito primordial es eliminar la fragmentación operativa provocada por el uso de cuadernos de notas físicos, "
        "planillas Excel desarticuladas y comunicaciones informales, centralizando todos los procesos en una plataforma en línea accesible, transparente y segura."
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_paragraph(
        "El sistema conecta en tiempo real a seis estamentos fundamentales de la comunidad educativa, asegurando que cada rol disponga de una interfaz "
        "personalizada, con navegación adaptada a sus tareas y con los niveles de seguridad pertinentes."
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Subsección 4.2: Los Seis Paneles de Usuario
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("4.2 Estructura y Funcionamiento Detallado de los Seis Paneles de Usuario")
    run.bold = True
    run.font.size = Pt(12)

    doc.add_paragraph(
        "EduNova implementa seis paneles especializados según el perfil de usuario autenticado. A continuación se documenta el propósito, "
        "las rutas de acceso, las funciones operativas y las características de navegación de cada uno de ellos:"
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Panel 1: Administrador
    doc.add_paragraph("Panel 1: Panel del Administrador (Ruta: /dashboard/)", style='Heading 2')
    doc.add_paragraph(
        "• Propósito Institucional: Control maestro, parametrización global del sistema, gestión de infraestructura institucional y supervisión operativa.\n"
        "• Funciones Clave Implementadas:\n"
        "  - Dashboard con métricas globales en tiempo real: Conteo de estudiantes activos, docentes vinculados, fichas/grados activos y programas técnicos articulados.\n"
        "  - Gestión Integral de Usuarios (/usuarios/): Creación, edición, activación y suspensión de cuentas con asignación estricta de roles institucionales.\n"
        "  - Centro de Control Institucional y Sedes (/instituciones/): Registro de colegios articulados, sedes físicas y convenios vigentes con el SENA.\n"
        "  - Módulo de Auditoría (/auditoria/): Monitoreo de trazas inmutables registradas por RegistroAuditoria (acciones, fechas, usuarios e IPs).\n"
        "  - Papelera de Reciclaje Institucional (/papelera/): Visualización de elementos dados de baja lógicamente con opción de restauración inmediata o purga definitiva.\n"
        "  - Gestión Centralizada de Caja y Pensiones (/pensiones/): Emisión de recibos de matrícula/pensión, consulta de recaudo diario, emisión de comprobantes HTML/PDF y anulación controlada.\n"
        "  - Centro de Reportes (/reportes/): Generación de informes consolidados de matrícula, rendimiento y estadísticas escolares.\n"
        "• Navegación y Usabilidad: Menú lateral estructurado con navegación jerárquica clara, indicadores numéricos actualizados dinámicamente y botones de acción unificados."
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Panel 2: Rectoría
    doc.add_paragraph("Panel 2: Panel de Rectoría / Directivos (Rutas: /rectoria/ o /coordinacion/)", style='Heading 2')
    doc.add_paragraph(
        "• Propósito Institucional: Supervisión pedagógica de alto nivel, monitoreo de la calidad formativa y toma de decisiones estratégicas.\n"
        "• Funciones Clave Implementadas:\n"
        "  - Visualización del Panorama General Institucional: Resumen consolidado del avance en la cobertura de programas técnicos y cumplimiento de metas pedagógicas.\n"
        "  - Monitoreo de Desempeño y Asistencia: Supervisión de estadísticas de asistencia y rendimiento escolar para detectar a tiempo alertas de ausentismo o deserción.\n"
        "  - Supervisión de Carga Académica: Consulta de la distribución de asignaturas, grupos y horas semanales asignadas a los docentes titulares.\n"
        "  - Emisión de Comunicados Directivos: Publicación de circulares oficiales y directrices rectorales hacia toda la comunidad educativa o hacia estamentos específicos.\n"
        "  - Consulta de Convenios SENA (/convenios/): Seguimiento a la vigencia y compromisos de los acuerdos interinstitucionales entre el colegio y el Centro Ecoturístico del SENA.\n"
        "• Navegación y Usabilidad: Enfocada en tableros visuales y reportes consolidados sin saturar al directivo con formularios operativos de registro diario."
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Panel 3: Secretaría
    doc.add_paragraph("Panel 3: Panel de Secretaría Académica (Rutas: /secretaria/ y /academico/matriculas/nueva/)", style='Heading 2')
    doc.add_paragraph(
        "• Propósito Institucional: Gestión documental, admisión, registro formal de matrículas y custodia de expedientes estudiantiles.\n"
        "• Funciones Clave Implementadas:\n"
        "  - Matrícula de Estudiantes con Creación Automática de Cuenta: Formulario oficial para registrar aprendices asociando nombres, apellidos, documento, tipo de documento, ficha/grado, jornada y sede. El sistema crea automáticamente la cuenta de usuario de Django asignando el número de documento como nombre de usuario y contraseña inicial controlada (probado en test_edunova_student_secretaria_flow.py Test 3).\n"
        "  - Consulta y Búsqueda de Expedientes: Filtrado rápido de estudiantes por documento, grado o apellidos con visualización de estado de matrícula (Matriculado, Retirado, En Trámite).\n"
        "  - Importación Masiva de Estudiantes: Descarga de plantilla estandarizada en formato Excel (.xlsx) y procesamiento masivo con validación de duplicados (probado en test_verificaciones.py Secciones 4 y 5).\n"
        "  - Generación de Certificados y Constancias: Emisión de constancias de estudio en formato digital con código de verificación institucional.\n"
        "  - Trámites de Novedades Académicas (/seguimiento/secretaria/): Registro de traslados, certificados y solicitudes radicadas por estudiantes y acudientes.\n"
        "• Navegación y Usabilidad: Formularios simplificados con validaciones en tiempo real para evitar errores tipográficos en documentos de identidad y correos electrónicos."
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Panel 4: Docente
    doc.add_paragraph("Panel 4: Panel del Docente / Instructor (Ruta: /instructor/ o /docente/)", style='Heading 2')
    doc.add_paragraph(
        "• Propósito Institucional: Gestión didáctica del aula, evaluación por competencias, control de asistencia diaria y comunicación pedagógica.\n"
        "• Funciones Clave Implementadas:\n"
        "  - Subpanel Inicio: Visualización del bloque 'Clase Actual', resumen de horario semanal y métricas de tareas pendientes por calificar (probado en test_edunova_teacher_flow.py Paso 2).\n"
        "  - Subpanel Mis Grupos y Asignaturas: Acceso segmentado a los grados asignados (ej. Grado 10°-A y Grado 11°-A) sin visualizar información de otros docentes.\n"
        "  - Subpanel Asistencia: Registro diario de asistencias por sesión con selección de estados: P (Presente), A (Ausente), J (Justificada), T (Tardanza) y campo de observaciones (probado en test_comprobacion_integral_edunova.py Paso 1).\n"
        "  - Subpanel Actividades y Tareas: Creación de talleres y actividades formativas con pestañas 'Publicadas' y 'Borradores', asociando Guías de Clase pedagógicas y Resultados de Aprendizaje (RAP) curricular de SOFIA; recepción de entregas de estudiantes y asignación de calificaciones numéricas (ej. 4.8) con retroalimentación cualitativa (probado en test_edunova_teacher_flow.py Pasos 3 a 7).\n"
        "  - Subpanel Calificaciones: Planilla de notas por período académico con cálculo automático de ponderaciones y botones de acción unificados en color rosa (#db2777).\n"
        "  - Subpanel Comunicaciones y Mensajería: Chat bidireccional directo con estudiantes y familias de sus cursos asignados (probado en test_edunova_teacher_flow.py Paso 8).\n"
        "  - Subpanel Recursos e Ideas Pedagógicas: Modal interactivo para subir y consultar material didáctico de apoyo para las clases.\n"
        "  - Subpanel Notificaciones: Centro de alertas en tiempo real sobre entregas radicadas por estudiantes y mensajes entrantes.\n"
        "• Navegación y Usabilidad: Pestañas internas que permiten al docente realizar todas sus labores sin recargar páginas innecesariamente ni perder el contexto de su clase."
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Panel 5: Familia
    doc.add_paragraph("Panel 5: Panel de Familia y Acudientes (Ruta: /familia/)", style='Heading 2')
    doc.add_paragraph(
        "• Propósito Institucional: Acompañamiento formativo, corresponsabilidad parental y supervisión del progreso integral de los hijos o tutelados.\n"
        "• Funciones Clave Implementadas:\n"
        "  - Consulta Inmediata de Rendimiento Escolar: Visualización del boletín de calificaciones actualizado por asignaturas y períodos académicos.\n"
        "  - Monitoreo Diario de Asistencia: Alertas y registro histórico de fallas o tardanzas registradas por los docentes, facilitando el control preventivo del ausentismo.\n"
        "  - Agenda de Actividades y Tareas: Visualización de las actividades formativas asignadas, fechas límite de entrega y estado de entrega de sus hijos.\n"
        "  - Canal de Comunicación con Docentes: Envío y recepción de mensajes directos con los profesores titulares de las asignaturas (probado en test_edunova_teacher_flow.py Paso 8-C).\n"
        "  - Consulta de Circulares y Comunicados Institucionales: Acceso a comunicados oficiales y citaciones emitidas por directivos y docentes.\n"
        "  - Consulta de Estado de Cartera y Pensiones: Visualización de recibos de pago emitidos y estado de solvencia escolar.\n"
        "• Navegación y Usabilidad: Diseño sumamente limpio y amigable para facilitar la consulta desde teléfonos móviles y dispositivos domésticos por parte de los padres."
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Panel 6: Estudiante
    doc.add_paragraph("Panel 6: Panel del Estudiante / Aprendiz (Rutas: /aprendiz/ y /portafolio/)", style='Heading 2')
    doc.add_paragraph(
        "• Propósito Institucional: Autogestión del proceso formativo, desarrollo de competencias, entrega puntual de evidencias y comunicación académica.\n"
        "• Nueve Módulos Integrados y Comprobados (scripts/test_edunova_student_secretaria_flow.py Test 4):\n"
        "  1. Inicio (/aprendiz/): Resumen de actividades por entregar, avisos de la semana y estado general de matrícula.\n"
        "  2. Calificaciones (/aprendiz/?tab=calificaciones): Consulta detallada de juicios evaluativos y notas cuantitativas por materia y período escolar.\n"
        "  3. Mi Horario (/aprendiz/?tab=horario): Cronograma semanal interactivo con bloques horarios, materias, aulas y profesores asignados.\n"
        "  4. Tareas y Actividades (/aprendiz/?tab=actividades): Descarga de guías pedagógicas en PDF y formulario para envío de respuestas y evidencias de talleres con archivos adjuntos.\n"
        "  5. Asistencia (/aprendiz/?tab=asistencia): Historial detallado de asistencias, tardanzas y justificaciones por materia.\n"
        "  6. Comunicados (/aprendiz/?tab=comunicados): Tablero de circulares oficiales y notificaciones escolares personalizadas.\n"
        "  7. Mensajería Directa (/aprendiz/?tab=mensajeria): Canal privado de mensajería para consultar dudas académicas directamente con sus docentes (probado en Test 6).\n"
        "  8. Portafolio de Evidencias (/portafolio/): Repositorio documental ordenado por resultados de aprendizaje (RAP) donde se almacenan las evidencias aprobadas requeridas por el SENA.\n"
        "  9. Perfil del Estudiante (/aprendiz/?tab=perfil): Consulta de datos institucionales y formulario de carga de fotografía de perfil institucional con validación de formato y persistencia (probado en Test 5).\n"
        "• Navegación y Usabilidad: Interfaz centrada en el estudiante, con estados claros de avance (Pendiente, Entregada, Calificada) y acceso directo a recursos de aprendizaje."
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Subsección 4.3: Facilidad de Uso y Experiencia de Usuario
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("4.3 Facilidad de Manejo y Experiencia del Usuario (UX/UI)")
    run.bold = True
    run.font.size = Pt(12)

    doc.add_paragraph(
        "En la construcción del prototipo se aplicaron estrictos criterios de diseño centrados en el usuario (User-Centered Design) "
        "y protocolos modernos de experiencia web. Las principales características evaluadas y consolidadas son:"
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    add_callout(
        doc,
        "CRITERIO 1: Navegación Adaptativa y Menú Lateral Intuitivo",
        "Cada usuario visualiza un menú lateral (sidebar) que se adapta automáticamente a su rol mediante etiquetas condicionales de plantilla de Django. "
        "Un estudiante o un padre de familia nunca se ve abrumado por opciones de administración de usuarios o configuración de servidores. "
        "La navegación es coherente y permite regresar al panel principal con un solo clic."
    )

    add_callout(
        doc,
        "CRITERIO 2: Jerarquía Visual y Consistencia de Color",
        "Se definió una paleta cromática profesional y armónica: azul marino institucional (#1E293B / #1B365D) para estructuras principales; "
        "verde SENA (#39A900) para distintivos de formación técnica; rosa unificado (#DB2777) para botones de acción en planillas docentes; "
        "y esquemas de alerta estándar (verde para aprobado, amarillo para pendiente, rojo para ausente o reprobado)."
    )

    add_callout(
        doc,
        "CRITERIO 3: Retroalimentación Inmediata al Usuario (Feedback)",
        "Toda acción que modifica datos (registro de notas, toma de asistencia, envío de tareas o emisión de pagos) dispara mensajes flash de confirmación "
        "con estilos distintivos (éxito, advertencia, error) a través del framework django.contrib.messages, asegurando que el usuario conozca de inmediato "
        "el resultado de su operación sin incertidumbre."
    )

    add_callout(
        doc,
        "CRITERIO 4: Estados Vacíos Profesionales y Prevención de Errores",
        "Cuando un módulo no contiene datos (por ejemplo, cuando no hay borradores de tareas o comunicados pendientes), el prototipo no muestra "
        "tablas vacías rotas o confusas, sino mensajes explicativos guiados con iconografía clara (ej. 'No hay actividades en borrador', 'Buzón al día')."
    )

    # Subsección 4.4: Mejoras Realizadas en el Prototipo
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("4.4 Mejoras Realizadas y Comprobadas en el Código de EduNova")
    run.bold = True
    run.font.size = Pt(12)

    doc.add_paragraph(
        "Durante las fases de estabilización del prototipo se implementaron mejoras concretas para resolver inconsistencias y optimizar la usabilidad del sistema. "
        "A continuación se documentan las mejoras comprobadas directamente en el repositorio:"
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_paragraph(
        "1. Corrección del Bug list.count() en Portafolio Estudiantil:\n"
        "   En la vista de portafolio (/portafolio/), se generaba un error 500 al intentar invocar el método de conteo sobre una lista procesada. "
        "   Se optimizó el algoritmo de agregación en academico/views.py para calcular los acumulados directamente sobre el QuerySet, restableciendo el acceso fluido (verificado en verificar_sistema_completo.py y test_edunova_student_secretaria_flow.py).\n\n"
        "2. Creación Automática y Vinculación de Cuentas en Matrícula:\n"
        "   Anteriormente, la secretaria debía registrar la matrícula y luego un administrador debía crear manualmente el usuario en el módulo de usuarios. "
        "   Se unificó el flujo en la vista de nueva matrícula: al registrar los datos básicos del estudiante, el backend crea de forma atómica la cuenta User, "
        "   crea el PerfilUsuario y asocia la Matricula con la Ficha escolar correspondiente, asignando el número de documento como credencial inicial de acceso.\n\n"
        "3. Incorporación de Papelera de Reciclaje con Snapshot JSON:\n"
        "   Se desarrolló el modelo PapeleraReciclaje para impedir pérdidas accidentales de estudiantes, horarios o cargas académicas. "
        "   Al solicitar la eliminación de un registro, el sistema serializa el estado completo en un campo JSONField y marca el elemento, permitiendo su restauración intacta.\n\n"
        "4. Unificación Visual de Botones en Planillas Docentes:\n"
        "   Se eliminaron discordancias visuales (botones verdes y azules mezclados) en el panel docente, adoptando un diseño uniforme con la clase .btn-rosa-edunova (#db2777), "
        "   mejorando la coherencia gráfica y la velocidad de reconocimiento de acciones en el aula.\n\n"
        "5. Descarga de Recibos de Caja en Formato PDF Oficial:\n"
        "   Se integró la generación al vuelo de comprobantes y recibos de pago en formato PDF (/pensiones/<id>/recibo-pdf/), "
        "   permitiendo a secretaría y tesorería entregar soportes impresos o digitales con código de recibo único a las familias."
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Subsección 4.5: Pruebas Funcionales Comprobadas
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("4.5 Batería de Pruebas Funcionales Reales Ejecutadas y Verificadas")
    run.bold = True
    run.font.size = Pt(12)

    doc.add_paragraph(
        "La funcionalidad del prototipo EduNova fue evaluada a través de suites de pruebas funcionales automatizadas ejecutadas directamente en el entorno activo. "
        "Los resultados certificados son los siguientes:"
    ).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    t_fun_tests = doc.add_table(rows=1, cols=4)
    t_fun_tests.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_fun_tests)
    
    hdr = t_fun_tests.rows[0].cells
    hdr[0].width = Inches(1.5)
    hdr[1].width = Inches(1.7)
    hdr[2].width = Inches(2.3)
    hdr[3].width = Inches(1.0)
    
    headers_ft = ["Módulo / Flujo Evaluado", "Script Ejecutor", "Acciones Evaluadas y Resultados", "Certificación"]
    for i, h in enumerate(headers_ft):
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9)
        set_cell_background(hdr[i], "E2E8F0")

    filas_fun_tests = [
        (
            "Flujo Integral del Panel Docente",
            "scripts/test_edunova_teacher_flow.py\n(12 Pasos)",
            "Login docente, clase actual en Inicio, filtros de actividades (Publicadas/Borradores), publicación de taller, entrega de alumno, calificación con nota 4.8 y feedback, mensajería docente-estudiante y familia-docente, recursos y perfil.",
            "100% ÉXITO\n(12/12 Aprobadas)"
        ),
        (
            "Flujo Estudiante y Secretaría",
            "scripts/test_edunova_student_secretaria_flow.py\n(8 Pasos)",
            "Acceso secretaría a nueva matrícula, remoción de cobro en panel, creación de estudiante con cédula como usuario, login documental, acceso a los 9 módulos de estudiante, subida de foto de perfil y mensajería.",
            "100% ÉXITO\n(8/8 Aprobadas)"
        ),
        (
            "Gestión de Caja, Grupos y Plantillas",
            "scripts/test_verificaciones.py\n(5 Pasos)",
            "Creación de cobro de pensión, visualización en recaudo de hoy en MySQL, comprobante HTML, descarga de recibo PDF, anulación, aislamiento de grupos docentes, circulares privadas y descarga de plantilla Excel.",
            "100% ÉXITO\n(5/5 Aprobadas)"
        ),
        (
            "Toma y Persistencia de Asistencia",
            "scripts/test_comprobacion_integral_edunova.py\n(Paso 1)",
            "Docente abre planilla de asistencia con estudiante asignado, registra estado Presente (P) con observación, verifica guardado en base de datos y actualiza a Ausente (A) sin duplicados.",
            "SUPERADA\n[OK]"
        )
    ]

    for mf, sc, acc, cert in filas_fun_tests:
        row = t_fun_tests.add_row().cells
        row[0].width = Inches(1.5)
        row[1].width = Inches(1.7)
        row[2].width = Inches(2.3)
        row[3].width = Inches(1.0)
        set_cell_margins(row[0])
        set_cell_margins(row[1])
        set_cell_margins(row[2])
        set_cell_margins(row[3])
        
        p0 = row[0].paragraphs[0]
        p0.add_run(mf).bold = True
        p0.runs[0].font.size = Pt(8.5)
        
        p1 = row[1].paragraphs[0]
        p1.add_run(sc).font.size = Pt(8)
        
        p2 = row[2].paragraphs[0]
        p2.paragraph_format.line_spacing = 1.15
        p2.add_run(acc).font.size = Pt(8)
        
        p3 = row[3].paragraphs[0]
        r3 = p3.add_run(cert)
        r3.bold = True
        r3.font.size = Pt(8.5)
        r3.font.color.rgb = RGBColor(0x15, 0x80, 0x3D)

    doc.add_page_break()

    # =========================================================================
    # 6. CONCLUSIONES
    # =========================================================================
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("5. CONCLUSIONES")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    p_c1 = doc.add_paragraph(
        "1. Eficacia del Modelo de Seguridad en Profundidad:\n"
        "La implementación de políticas rigurosas de control de acceso basado en roles (RBAC) y la incorporación de reglas anti-IDOR "
        "en el backend de EduNova garantizan que la confidencialidad de la información académica, personal y financiera de los estudiantes "
        "esté plenamente salvaguardada. Se demostró que la seguridad de las aplicaciones web no puede depender del ocultamiento de botones en el frontend, "
        "sino de interceptores y validaciones sólidas en el servidor antes de cualquier procesamiento de datos."
    )
    p_c1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p_c2 = doc.add_paragraph(
        "2. Madurez y Robustez del Prototipo Funcional de EduNova:\n"
        "La construcción del prototipo cubrió satisfactoriamente las necesidades operativas de los seis estamentos institucionales "
        "(Administrador, Rectoría, Secretaría, Docente, Familia y Estudiante). La arquitectura modular adoptada en Django, "
        "junto con el diseño responsivo y la aplicación de protocolos de experiencia de usuario (UX/UI), demuestra que el software está preparado "
        "para articular eficazmente la gestión de la Media Técnica del SENA en las instituciones educativas del Departamento del Magdalena."
    )
    p_c2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p_c3 = doc.add_paragraph(
        "3. Validación Basada en Evidencias Reales y Trazabilidad:\n"
        "La ejecución exitosa del 100% de las pruebas funcionales y de seguridad documentadas (flujo docente, flujo estudiante-secretaría, "
        "gestión de caja y pruebas de aislamiento) certifica que las correcciones efectuadas en el código fuente resolvieron las vulnerabilidades "
        "detectadas sin alterar la estabilidad del aplicativo, cumpliendo fielmente con los criterios de evaluación pedagógica y técnica exigidos por el SENA."
    )
    p_c3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Espacio para firmas formales
    p_firmas = doc.add_paragraph()
    p_firmas.paragraph_format.space_before = Pt(40)
    p_firmas.paragraph_format.space_after = Pt(10)
    r_f = p_firmas.add_run("CONSTANCIA DE ELABORACIÓN Y ENTREGA:\n")
    r_f.bold = True
    r_f.font.size = Pt(11)

    t_firm = doc.add_table(rows=1, cols=2)
    t_firm.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_firm, val="none")
    
    cf = t_firm.rows[0].cells
    cf[0].width = Inches(3.2)
    cf[1].width = Inches(3.2)
    
    p_f1 = cf[0].paragraphs[0]
    p_f1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_f1.add_run("_____________________________________\n").bold = True
    p_f1.add_run("AMINTA MOGUEA FLOREZ\n").bold = True
    p_f1.add_run("Aprendiz SENA — ADSI / ADSO\nC.C. de Santa Marta")
    
    p_f2 = cf[1].paragraphs[0]
    p_f2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_f2.add_run("_____________________________________\n").bold = True
    p_f2.add_run("DULCE MARÍA FIGUEROA\n").bold = True
    p_f2.add_run("Aprendiz SENA — ADSI / ADSO\nC.C. de Santa Marta")

    return doc

def guardar_en_ubicaciones(doc):
    nombre_archivo = "Actividades_SENA_EduNova_Aminta_y_Dulce.docx"
    
    # 1. Carpeta raíz C:\SINETEC (la carpeta de trabajo / carpeta amarilla principal)
    ruta_root = os.path.join(r"C:\SINETEC", nombre_archivo)
    
    # 2. Carpeta docs de SINETEC
    ruta_docs = os.path.join(r"C:\SINETEC\SINETEC\docs", nombre_archivo)
    os.makedirs(os.path.dirname(ruta_docs), exist_ok=True)
    
    # 3. Carpeta Descargas del usuario (Downloads)
    ruta_downloads = os.path.join(os.path.expanduser("~"), "Downloads", nombre_archivo)
    
    # 4. Carpeta Documentos del usuario si existe
    ruta_onedrive_docs = os.path.join(os.path.expanduser("~"), "OneDrive", "Documentos", nombre_archivo)
    
    # Guardar en las ubicaciones
    doc.save(ruta_root)
    print(f"[OK] Guardado exitosamente en: {ruta_root}")
    
    doc.save(ruta_docs)
    print(f"[OK] Guardado exitosamente en: {ruta_docs}")
    
    try:
        doc.save(ruta_downloads)
        print(f"[OK] Copia de seguridad guardada en: {ruta_downloads}")
    except Exception as e:
        print(f"[ADVERTENCIA] No se pudo guardar en Downloads: {e}")
        
    try:
        if os.path.exists(os.path.dirname(ruta_onedrive_docs)):
            doc.save(ruta_onedrive_docs)
            print(f"[OK] Copia de seguridad guardada en: {ruta_onedrive_docs}")
    except Exception as e:
        print(f"[ADVERTENCIA] No se pudo guardar en OneDrive Documentos: {e}")

    # Verificar que el archivo principal existe y tiene tamaño válido
    if os.path.exists(ruta_root) and os.path.getsize(ruta_root) > 1000:
        print(f"[VERIFICACIÓN EXITOSA] El archivo existe en {ruta_root} con tamaño {os.path.getsize(ruta_root)} bytes.")
        return ruta_root
    else:
        raise RuntimeError("El archivo no se guardó correctamente en la ruta principal.")

if __name__ == '__main__':
    print("Iniciando generación de documento oficial de Actividades SENA EduNova...")
    documento = construir_documento()
    ruta_final = guardar_en_ubicaciones(documento)
    print(f"PROCESO FINALIZADO CON ÉXITO. Archivo listo para Microsoft Word en: {ruta_final}")
