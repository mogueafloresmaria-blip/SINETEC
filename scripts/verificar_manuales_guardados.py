# -*- coding: utf-8 -*-
import os
import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

rutas = [
    r"C:\EDUNOVA\EDUNOVA\docs\Manual_de_Usuario_EduNova.docx",
    r"C:\EDUNOVA\EDUNOVA\docs\Manual_de_Usuario_EduNova.pdf",
    r"C:\EDUNOVA\EDUNOVA\docs\Manual_Tecnico_EduNova.docx",
    r"C:\EDUNOVA\EDUNOVA\docs\Manual_Tecnico_EduNova.pdf",
    r"C:\EDUNOVA\docs\Manual_de_Usuario_EduNova.docx",
    r"C:\EDUNOVA\docs\Manual_Tecnico_EduNova.docx",
    r"C:\EDUNOVA\EDUNOVA\docs\09_manuales\Manual_de_Usuario_EduNova.docx",
    r"C:\EDUNOVA\EDUNOVA\docs\09_manuales\Manual_Tecnico_EduNova.docx"
]

print("=" * 70)
print("VERIFICACIÓN DE ARCHIVOS FÍSICOS EN DISCO LOCAL")
print("=" * 70)

for r in rutas:
    if os.path.exists(r):
        sz = os.path.getsize(r)
        print(f"EXISTE: {r} ({sz:,} bytes)")
        if r.endswith(".docx"):
            d = docx.Document(r)
            chars = sum(len(p.text) for p in d.paragraphs)
            print(f"   -> Párrafos: {len(d.paragraphs)}, Tablas: {len(d.tables)}, Caracteres: {chars:,}")
    else:
        print(f"NO EXISTE: {r}")

print("=" * 70)
