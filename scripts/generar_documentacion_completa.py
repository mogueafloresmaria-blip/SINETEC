# -*- coding: utf-8 -*-
import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

# Ensure workspace root is in path
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

# Import data
from scripts.generar_documentos_md_html import (
    TITULO, SUBTITULO, INSTITUCION, SECCIONES_DATA, DOCS_DIR
)

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F5EFE6")
    set_cell_margins(cell, top=120, bottom=120, left=200, right=200)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x3D, 0x23, 0x14)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    if level == 1:
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        for r in h.runs:
            r.font.name = 'Calibri'
            r.font.size = Pt(15)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x3D, 0x23, 0x14)
    elif level == 2:
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        for r in h.runs:
            r.font.name = 'Calibri'
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x8C, 0x58, 0x35)
    return h

def build_docx_documentation(docx_path):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    # PORTADA
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(36)
    p_pre.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_pre.add_run(INSTITUCION)
    r_inst.font.name = 'Calibri'
    r_inst.font.size = Pt(12)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(0x8C, 0x58, 0x35)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(20)
    p_title.paragraph_format.space_after = Pt(10)
    p_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(TITULO.replace(': ', ':\n'))
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x3D, 0x23, 0x14)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(36)
    p_sub.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run(SUBTITULO)
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x5A, 0x3E, 0x2B)

    # Meta Table on Cover
    meta_tbl = doc.add_table(rows=6, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Proyecto:", "EduNova (SINETEC)"),
        ("Tipo de Documento:", "Manual T\u00e9cnico de Despliegue en Producci\u00f3n"),
        ("Destinatario / Evaluaci\u00f3n:", "Presentaci\u00f3n Docente / Entrega de Proyecto"),
        ("Versi\u00f3n del Sistema:", "2.4.0 (Listo para Producci\u00f3n)"),
        ("Pila Tecnol\u00f3gica:", "Django 5.x | Python 3.13 | MySQL 8.0 | Gunicorn | Nginx"),
        ("Fecha:", "Octubre de 2026")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_tbl.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        set_cell_background(c0, "FAEDCD")
        set_cell_background(c1, "FDFBF7")
        set_cell_margins(c0, top=80, bottom=80, left=120, right=120)
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = RGBColor(0x3D, 0x23, 0x14)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    doc.add_page_break()

    # ÍNDICE
    add_styled_heading(doc, "\u00cdNDICE GENERAL DE CONTENIDOS", 1)
    for sec in SECCIONES_DATA:
        p_toc = doc.add_paragraph()
        p_toc.paragraph_format.space_before = Pt(2)
        p_toc.paragraph_format.space_after = Pt(2)
        p_toc.paragraph_format.line_spacing = 1.15
        r_toc = p_toc.add_run(f"{sec['num']}. {sec['title']}")
        r_toc.font.size = Pt(10)
        r_toc.font.bold = True
        r_toc.font.color.rgb = RGBColor(0x3D, 0x23, 0x14)

    doc.add_page_break()

    # SECCIONES
    for sec in SECCIONES_DATA:
        add_styled_heading(doc, f"{sec['num']}. {sec['title']}", 1)
        paragraphs = sec['content'].split('\n\n')
        for par in paragraphs:
            if any(par.strip().startswith(cmd) for cmd in [
                'sudo ', 'pip ', 'python ', 'git ', 'chmod ', 'chown ', 'cd ',
                '[Unit]', 'server {', '#!/bin/bash', 'CREATE ', '-- '
            ]):
                add_code_block(doc, par)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                r = p.add_run(par.strip())
                r.font.name = 'Calibri'
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    doc.save(docx_path)
    print(f"[OK] Documento Word regenerado en: {docx_path}")

if __name__ == "__main__":
    target_docx = os.path.join(DOCS_DIR, "Documentacion_Tecnica_Instalacion_Despliegue_EduNova.docx")
    build_docx_documentation(target_docx)
