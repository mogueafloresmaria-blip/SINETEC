# -*- coding: utf-8 -*-
"""
Script de Renderizado Dual (DOCX y PDF) para el Informe Administrativo de EduNova para el SENA.
Genera los archivos:
- docs/Informe_Administrativo_EduNova_SENA.docx
- docs/Informe_Administrativo_EduNova_SENA.pdf
"""

import os
import sys

# Ajustar PYTHONPATH
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from scripts.informe_content import (
    METADATOS, CONTROL_VERSIONES, INTRODUCCION_TEXT, OBJETIVOS_DATA,
    ALCANCE_DATA, MATRIZ_ESTADO, REQUISITOS_SERVIDOR, PLAN_INSTALACION_FASES,
    VARIABLES_ENTORNO_DATA, PLAN_RESPALDO_DATA, PLAN_MIGRACION_DATA,
    PRUEBAS_EVIDENCIAS_DATA, SEGURIDAD_DATA, MANTENIMIENTO_DATA,
    RIESGOS_DATA, CONCLUSIONES_DATA, RECOMENDACIONES_DATA, FIRMAS_DATA
)

DOCS_DIR = os.path.join(BASE_DIR, "docs")
os.makedirs(DOCS_DIR, exist_ok=True)
DOCX_PATH = os.path.join(DOCS_DIR, "Informe_Administrativo_EduNova_SENA.docx")
PDF_PATH = os.path.join(DOCS_DIR, "Informe_Administrativo_EduNova_SENA.pdf")

# ==============================================================================
# IMPORTS DOCX Y REPORTLAB
# ==============================================================================
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# Registrar fuentes para ReportLab
pdfmetrics.registerFont(TTFont('Arial', 'C:/Windows/Fonts/arial.ttf'))
pdfmetrics.registerFont(TTFont('Arial-Bold', 'C:/Windows/Fonts/arialbd.ttf'))
pdfmetrics.registerFont(TTFont('Arial-Italic', 'C:/Windows/Fonts/ariali.ttf'))
pdfmetrics.registerFont(TTFont('Consolas', 'C:/Windows/Fonts/consola.ttf'))

# ==============================================================================
# CLASE CANVAS NUMERADO PARA REPORTLAB
# ==============================================================================
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_decorations(self, page_count):
        if self._pageNumber == 1:
            return  # Portada sin encabezado ni pie de página
        
        self.saveState()
        self.setFont("Arial", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        
        # Encabezado superior
        self.drawString(54, 11 * inch - 36, "SENA | Regional Magdalena — Informe Administrativo de Despliegue, Respaldo y Migración")
        self.drawRightString(8.5 * inch - 54, 11 * inch - 36, "Sistema EduNova")
        self.setStrokeColor(colors.HexColor("#39A900"))  # Verde SENA
        self.setLineWidth(0.75)
        self.line(54, 11 * inch - 40, 8.5 * inch - 54, 11 * inch - 40)
        
        # Pie de página inferior
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(54, 46, 8.5 * inch - 54, 46)
        self.drawString(54, 32, "Centro de Logística y Promoción Ecoturística del Magdalena • Ficha 228106")
        self.drawRightString(8.5 * inch - 54, 32, f"Página {self._pageNumber} de {page_count}")
        self.restoreState()

# ==============================================================================
# HELPERS DOCX
# ==============================================================================
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

# ==============================================================================
# CONSTRUCTOR DOCX
# ==============================================================================
def generar_documento_docx(ruta_salida):
    print("Iniciando construcción de archivo DOCX...")
    doc = docx.Document()
    
    # Configuración de márgenes
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Encabezado
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.text = "SENA REGIONAL MAGDALENA • INFORME ADMINISTRATIVO DE DESPLIEGUE, RESPALDO Y MIGRACIÓN"
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_hdr.style.font.name = "Arial"
        p_hdr.style.font.size = Pt(8.5)
        p_hdr.style.font.color.rgb = RGBColor(100, 116, 139)
        
        # Pie de página
        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.text = "Sistema de Información EduNova (SINETEC) • Ficha 228106 • Documento Oficial SENA"
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ftr.style.font.name = "Arial"
        p_ftr.style.font.size = Pt(8.5)
        p_ftr.style.font.color.rgb = RGBColor(100, 116, 139)

    # PORTADA
    p_badge = doc.add_paragraph()
    p_badge.paragraph_format.space_before = Pt(20)
    p_badge.paragraph_format.space_after = Pt(8)
    run_badge = p_badge.add_run("SERVICIO NACIONAL DE APRENDIZAJE – SENA | REGIONAL MAGDALENA")
    run_badge.font.name = "Arial"
    run_badge.font.size = Pt(11)
    run_badge.font.bold = True
    run_badge.font.color.rgb = RGBColor(57, 169, 0) # Verde SENA

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(8)
    run_title = p_title.add_run("INFORME ADMINISTRATIVO DE INSTALACIÓN, RESPALDO, MIGRACIÓN Y DESPLIEGUE A PRODUCCIÓN")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 41, 59) # Slate

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run("SISTEMA DE INFORMACIÓN EDUNOVA — SISTEMA INTEGRADO DE EDUCACIÓN TÉCNICA (SINETEC)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(20)
    run_div = p_div.add_run("―" * 58)
    run_div.font.color.rgb = RGBColor(57, 169, 0)

    # Tabla de metadatos en portada
    tbl_meta = doc.add_table(rows=8, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_meta, color="CBD5E1")
    meta_info = [
        ("Programa de Formación:", METADATOS["programa"]),
        ("Código de Ficha / Proyecto:", f"Ficha {METADATOS['ficha']} | Proyecto SOFIA: {METADATOS['proyecto_sofia']}"),
        ("Centro de Formación:", METADATOS["centro"]),
        ("Nombre del Proyecto:", METADATOS["nombre_proyecto"]),
        ("Versión del Documento:", METADATOS["version_documento"]),
        ("Destinatario / Evaluación:", METADATOS["destinatario"]),
        ("Autor / Desarrollador:", METADATOS["autor"]),
        ("Ciudad y Fecha:", f"{METADATOS['ciudad']} — {METADATOS['fecha']}")
    ]
    for idx, (lbl, val) in enumerate(meta_info):
        row = tbl_meta.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.3)
        c1.width = Inches(4.2)
        c0.paragraphs[0].text = lbl
        c0.paragraphs[0].runs[0].font.name = "Arial"
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.size = Pt(9.5)
        c0.paragraphs[0].runs[0].font.color.rgb = RGBColor(30, 41, 59)
        c1.paragraphs[0].text = val
        c1.paragraphs[0].runs[0].font.name = "Arial"
        c1.paragraphs[0].runs[0].font.size = Pt(9.5)
        c1.paragraphs[0].runs[0].font.color.rgb = RGBColor(51, 65, 85)
        bg = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_margins(c0, top=90, bottom=90, left=120, right=120)
        set_cell_margins(c1, top=90, bottom=90, left=120, right=120)

    doc.add_page_break()

    # FUNCIONES AUXILIARES PARA EL CONTENIDO
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        run = h.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(57, 169, 0) # Verde SENA
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        run = h.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59)
        return h

    def add_h3(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)
        run = h.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(71, 85, 105)
        return h

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Arial"
            r_pre.font.size = Pt(10)
            r_pre.font.bold = True
            r_pre.font.color.rgb = RGBColor(30, 41, 59)
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def add_alert(title, text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        cell.width = Inches(6.5)
        set_cell_background(cell, "F0FDF4") # Fondo verde muy claro
        set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
        
        # Borde izquierdo verde
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="39A900"/>'
            f'  <w:top w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'  <w:bottom w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r_t = p.add_run(f"NOTA INSTITUCIONAL • {title}\n")
        r_t.font.name = "Arial"
        r_t.font.size = Pt(9.5)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(22, 101, 52)
        
        r_b = p.add_run(text)
        r_b.font.name = "Arial"
        r_b.font.size = Pt(9)
        r_b.font.color.rgb = RGBColor(22, 101, 52)
        
        p_sep = doc.add_paragraph()
        p_sep.paragraph_format.space_before = Pt(0)
        p_sep.paragraph_format.space_after = Pt(4)

    def add_code(lines):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        cell.width = Inches(6.5)
        set_cell_background(cell, "0F172A") # Slate oscuro
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run("\n".join(lines))
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(226, 232, 240)
        
        p_sep = doc.add_paragraph()
        p_sep.paragraph_format.space_before = Pt(0)
        p_sep.paragraph_format.space_after = Pt(4)

    # --------------------------------------------------------------------------
    # SECCIÓN: CONTROL DE VERSIONES
    # --------------------------------------------------------------------------
    add_h1("CONTROL DE VERSIONES Y REGISTRO DOCUMENTAL")
    add_p("El presente informe se rige por el estándar de trazabilidad y gestión de configuración documental del SENA:")
    tbl_ver = doc.add_table(rows=len(CONTROL_VERSIONES)+1, cols=4)
    tbl_ver.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_ver, color="CBD5E1")
    headers_ver = ["Versión", "Fecha", "Autor", "Descripción del Cambio / Hito"]
    col_w_ver = [Inches(0.9), Inches(1.1), Inches(1.7), Inches(2.8)]
    for j, h in enumerate(headers_ver):
        cell = tbl_ver.rows[0].cells[j]
        cell.width = col_w_ver[j]
        cell.paragraphs[0].text = h
        cell.paragraphs[0].runs[0].font.name = "Arial"
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
    for i, item in enumerate(CONTROL_VERSIONES):
        row = tbl_ver.rows[i+1]
        vals = [item["version"], item["fecha"], item["autor"], item["descripcion"]]
        bg = "F8FAFC" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate(vals):
            cell = row.cells[j]
            cell.width = col_w_ver[j]
            cell.paragraphs[0].text = val
            cell.paragraphs[0].runs[0].font.name = "Arial"
            cell.paragraphs[0].runs[0].font.size = Pt(8.5)
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(51, 65, 85)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)

    # --------------------------------------------------------------------------
    # 1. INTRODUCCIÓN
    # --------------------------------------------------------------------------
    add_h1("1. INTRODUCCIÓN Y CONTEXTO ADMINISTRATIVO")
    for par in INTRODUCCION_TEXT.split("\n\n"):
        add_p(par)

    add_alert(
        "DISTINCIÓN DE ESTADOS EN EL DOCUMENTO",
        "Conforme a lo instruido por la coordinación del SENA, este informe no simula servidores o dominios de producción ficticios. Se certifica que EduNova se encuentra desarrollado, estabilizado y verificado localmente con suites automatizadas (100% de éxito), y los procedimientos de producción se detallan como propuesta formal estandarizada pendiente de despliegue sobre el servidor final provisto por la institución educativa."
    )

    # --------------------------------------------------------------------------
    # 2. OBJETIVOS
    # --------------------------------------------------------------------------
    add_h1("2. OBJETIVOS DEL INFORME")
    add_h2("2.1 Objetivo General")
    add_p(OBJETIVOS_DATA["general"])
    add_h2("2.2 Objetivos Específicos")
    for obj in OBJETIVOS_DATA["especificos"]:
        add_p(obj, bold_prefix="• ")

    # --------------------------------------------------------------------------
    # 3. ALCANCE Y LÍMITES
    # --------------------------------------------------------------------------
    add_h1("3. ALCANCE DEL SISTEMA Y LÍMITES DEL DESPLIEGUE")
    add_h2("3.1 Alcance Funcional de la Solución (Seis Paneles)")
    for p_fun in ALCANCE_DATA["funcional"]:
        parts = p_fun.split(": ")
        add_p(parts[1], bold_prefix=f"• {parts[0]}: ")
    add_h2("3.2 Alcance del Despliegue en Servidor")
    add_p(ALCANCE_DATA["despliegue"])
    add_h2("3.3 Límites y Exclusiones del Sistema")
    for lim in ALCANCE_DATA["limites"]:
        add_p(lim, bold_prefix="• Exclusión: ")

    # --------------------------------------------------------------------------
    # 4. MATRIZ DE ESTADO
    # --------------------------------------------------------------------------
    add_h1("4. MATRIZ DE ESTADO REAL DEL PROYECTO")
    add_p("A continuación se consolida el estado verificado de cada componente técnico, distinguiendo entre lo realizado, propuesto y pendiente:")
    tbl_est = doc.add_table(rows=len(MATRIZ_ESTADO)+1, cols=4)
    tbl_est.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_est, color="CBD5E1")
    headers_est = ["Área Técnica", "Actividad / Componente", "Estado Formal", "Evidencia Verificable"]
    col_w_est = [Inches(1.3), Inches(2.2), Inches(1.1), Inches(1.9)]
    for j, h in enumerate(headers_est):
        cell = tbl_est.rows[0].cells[j]
        cell.width = col_w_est[j]
        cell.paragraphs[0].text = h
        cell.paragraphs[0].runs[0].font.name = "Arial"
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=80, bottom=80, left=90, right=90)
    for i, item in enumerate(MATRIZ_ESTADO):
        row = tbl_est.rows[i+1]
        vals = [item["area"], item["actividad"], item["tipo"], item["evidencia"]]
        bg = "F8FAFC" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate(vals):
            cell = row.cells[j]
            cell.width = col_w_est[j]
            cell.paragraphs[0].text = val
            cell.paragraphs[0].runs[0].font.name = "Arial"
            cell.paragraphs[0].runs[0].font.size = Pt(8)
            if j == 2:
                cell.paragraphs[0].runs[0].font.bold = True
                if val == "REALIZADA":
                    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(22, 101, 52)
                elif val == "PROPUESTA":
                    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(30, 64, 175)
                else:
                    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(180, 83, 9)
            else:
                cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(51, 65, 85)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)

    # --------------------------------------------------------------------------
    # 5. REQUISITOS DEL SERVIDOR
    # --------------------------------------------------------------------------
    add_h1("5. REQUISITOS DEL SERVIDOR (INFRAESTRUCTURA, HARDWARE, SOFTWARE Y RED)")
    add_h2("5.1 Requisitos de Hardware del Servidor")
    hw_rows = REQUISITOS_SERVIDOR["hardware"]
    tbl_hw = doc.add_table(rows=len(hw_rows), cols=4)
    tbl_hw.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_hw, color="CBD5E1")
    col_w_hw = [Inches(1.2), Inches(1.5), Inches(1.7), Inches(2.1)]
    for i, row_data in enumerate(hw_rows):
        row = tbl_hw.rows[i]
        is_hdr = (i == 0)
        bg = "1E293B" if is_hdr else ("F8FAFC" if i % 2 == 1 else "FFFFFF")
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            cell.width = col_w_hw[j]
            cell.paragraphs[0].text = val
            cell.paragraphs[0].runs[0].font.name = "Arial"
            cell.paragraphs[0].runs[0].font.size = Pt(8.5) if not is_hdr else Pt(9)
            cell.paragraphs[0].runs[0].font.bold = is_hdr or (j == 0)
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255) if is_hdr else RGBColor(51, 65, 85)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)

    add_h2("5.2 Requisitos de Software Homologado")
    sw_rows = REQUISITOS_SERVIDOR["software"]
    tbl_sw = doc.add_table(rows=len(sw_rows), cols=4)
    tbl_sw.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_sw, color="CBD5E1")
    col_w_sw = [Inches(1.3), Inches(1.4), Inches(1.4), Inches(2.4)]
    for i, row_data in enumerate(sw_rows):
        row = tbl_sw.rows[i]
        is_hdr = (i == 0)
        bg = "1E293B" if is_hdr else ("F8FAFC" if i % 2 == 1 else "FFFFFF")
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            cell.width = col_w_sw[j]
            cell.paragraphs[0].text = val
            cell.paragraphs[0].runs[0].font.name = "Arial"
            cell.paragraphs[0].runs[0].font.size = Pt(8.5) if not is_hdr else Pt(9)
            cell.paragraphs[0].runs[0].font.bold = is_hdr or (j == 0)
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255) if is_hdr else RGBColor(51, 65, 85)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)

    add_h2("5.3 Requisitos de Red, Puertos y Firewall")
    net_rows = REQUISITOS_SERVIDOR["red_puertos"]
    tbl_net = doc.add_table(rows=len(net_rows), cols=5)
    tbl_net.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_net, color="CBD5E1")
    col_w_net = [Inches(0.6), Inches(0.8), Inches(1.3), Inches(1.3), Inches(2.5)]
    for i, row_data in enumerate(net_rows):
        row = tbl_net.rows[i]
        is_hdr = (i == 0)
        bg = "1E293B" if is_hdr else ("F8FAFC" if i % 2 == 1 else "FFFFFF")
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            cell.width = col_w_net[j]
            cell.paragraphs[0].text = val
            cell.paragraphs[0].runs[0].font.name = "Arial"
            cell.paragraphs[0].runs[0].font.size = Pt(8) if not is_hdr else Pt(8.5)
            cell.paragraphs[0].runs[0].font.bold = is_hdr or (j == 0)
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255) if is_hdr else RGBColor(51, 65, 85)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)

    # --------------------------------------------------------------------------
    # 6. PLAN DETALLADO DE INSTALACIÓN Y DESPLIEGUE
    # --------------------------------------------------------------------------
    add_h1("6. PLAN DETALLADO DE INSTALACIÓN Y DESPLIEGUE A PRODUCCIÓN")
    add_p("A continuación se describe el procedimiento paso a paso para la puesta en marcha técnica de EduNova sobre un servidor institucional dedicado con Ubuntu Server:")
    for fase in PLAN_INSTALACION_FASES:
        add_h2(fase["fase"])
        add_p(fase["descripcion"])
        add_code(fase["comandos"])

    # --------------------------------------------------------------------------
    # 7. CONFIGURACIÓN DEL SISTEMA Y VARIABLES DE ENTORNO
    # --------------------------------------------------------------------------
    add_h1("7. CONFIGURACIÓN DEL SISTEMA Y VARIABLES DE ENTORNO")
    add_p("El sistema utiliza el paquete python-decouple para aislar secretos en un archivo .env con permisos 600, previniendo fuga de claves:")
    tbl_env = doc.add_table(rows=len(VARIABLES_ENTORNO_DATA)+1, cols=4)
    tbl_env.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_env, color="CBD5E1")
    headers_env = ["Variable de Entorno", "Propósito en el Sistema", "Ejemplo en Producción", "Criticidad"]
    col_w_env = [Inches(1.5), Inches(2.2), Inches(1.8), Inches(1.0)]
    for j, h in enumerate(headers_env):
        cell = tbl_env.rows[0].cells[j]
        cell.width = col_w_env[j]
        cell.paragraphs[0].text = h
        cell.paragraphs[0].runs[0].font.name = "Arial"
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(8.5)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
    for i, row_data in enumerate(VARIABLES_ENTORNO_DATA):
        row = tbl_env.rows[i+1]
        bg = "F8FAFC" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            cell.width = col_w_env[j]
            cell.paragraphs[0].text = val
            cell.paragraphs[0].runs[0].font.name = "Arial"
            cell.paragraphs[0].runs[0].font.size = Pt(8)
            cell.paragraphs[0].runs[0].font.bold = (j == 0)
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(51, 65, 85)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)

    # --------------------------------------------------------------------------
    # 8. PLAN DE RESPALDO Y RECUPERACIÓN
    # --------------------------------------------------------------------------
    add_h1("8. PLAN DE RESPALDO Y RECUPERACIÓN (BACKUP & DISASTER RECOVERY)")
    add_p(PLAN_RESPALDO_DATA["descripcion"])
    add_h2("8.1 Política y Periodicidad de Respaldos")
    tbl_pol = doc.add_table(rows=len(PLAN_RESPALDO_DATA["politica"]), cols=5)
    tbl_pol.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_pol, color="CBD5E1")
    col_w_pol = [Inches(1.3), Inches(1.2), Inches(1.3), Inches(1.2), Inches(1.5)]
    for i, row_data in enumerate(PLAN_RESPALDO_DATA["politica"]):
        row = tbl_pol.rows[i]
        is_hdr = (i == 0)
        bg = "1E293B" if is_hdr else ("F8FAFC" if i % 2 == 1 else "FFFFFF")
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            cell.width = col_w_pol[j]
            cell.paragraphs[0].text = val
            cell.paragraphs[0].runs[0].font.name = "Arial"
            cell.paragraphs[0].runs[0].font.size = Pt(8)
            cell.paragraphs[0].runs[0].font.bold = is_hdr or (j == 0)
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255) if is_hdr else RGBColor(51, 65, 85)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)

    add_h2("8.2 Script de Respaldo Automatizado (/usr/local/bin/backup_edunova.sh)")
    add_code(PLAN_RESPALDO_DATA["script_bash"])
    add_p(f"Programación en cron del servidor (ejecución automática a las 02:00 AM):", bold_prefix="Entrada Crontab: ")
    add_code([PLAN_RESPALDO_DATA["cron_entry"]])

    add_h2("8.3 Procedimiento de Restauración ante Desastres (Disaster Recovery)")
    for paso in PLAN_RESPALDO_DATA["restauracion_pasos"]:
        add_p(paso)

    add_h2("8.4 Inventario de Respaldos Reales Existentes en el Repositorio")
    tbl_bak = doc.add_table(rows=len(PLAN_RESPALDO_DATA["backups_reales"])+1, cols=4)
    tbl_bak.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_bak, color="CBD5E1")
    col_w_bak = [Inches(1.8), Inches(0.8), Inches(1.8), Inches(2.1)]
    headers_bak = ["Archivo de Respaldo", "Tamaño", "Ruta en Repositorio", "Descripción del Contenido"]
    for j, h in enumerate(headers_bak):
        cell = tbl_bak.rows[0].cells[j]
        cell.width = col_w_bak[j]
        cell.paragraphs[0].text = h
        cell.paragraphs[0].runs[0].font.name = "Arial"
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(8.5)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
    for i, row_data in enumerate(PLAN_RESPALDO_DATA["backups_reales"]):
        row = tbl_bak.rows[i+1]
        bg = "F8FAFC" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            cell.width = col_w_bak[j]
            cell.paragraphs[0].text = val
            cell.paragraphs[0].runs[0].font.name = "Arial"
            cell.paragraphs[0].runs[0].font.size = Pt(8)
            cell.paragraphs[0].runs[0].font.bold = (j == 0)
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(51, 65, 85)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)

    # --------------------------------------------------------------------------
    # 9. PLAN DE MIGRACIÓN DE DATOS
    # --------------------------------------------------------------------------
    add_h1("9. PLAN DE MIGRACIÓN DE DATOS")
    add_p(PLAN_MIGRACION_DATA["metodologia"])
    add_h2("9.1 Etapas del Proceso de Migración (ETL)")
    for tit, desc in PLAN_MIGRACION_DATA["etapas"]:
        add_p(desc, bold_prefix=f"{tit}: ")
    add_h2("9.2 Chequeos y Verificaciones Post-Migración")
    for chk in PLAN_MIGRACION_DATA["chequeos_post_migracion"]:
        add_p(chk, bold_prefix="• SQL / Verificación: ")

    # --------------------------------------------------------------------------
    # 10. PRUEBAS DE FUNCIONAMIENTO Y CONTROL DE CALIDAD
    # --------------------------------------------------------------------------
    add_h1("10. PRUEBAS DE FUNCIONAMIENTO Y CONTROL DE CALIDAD")
    add_p("El sistema fue sometido a una suite exhaustiva de pruebas unitarias y de integración que validan el 100% de las reglas de negocio institucionales:")
    tbl_pru = doc.add_table(rows=len(PRUEBAS_EVIDENCIAS_DATA)+1, cols=5)
    tbl_pru.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_pru, color="CBD5E1")
    col_w_pru = [Inches(0.8), Inches(1.4), Inches(1.8), Inches(1.0), Inches(1.5)]
    headers_pru = ["ID", "Nombre de Prueba", "Resultado Esperado vs Obtenido", "Estado", "Evidencia"]
    for j, h in enumerate(headers_pru):
        cell = tbl_pru.rows[0].cells[j]
        cell.width = col_w_pru[j]
        cell.paragraphs[0].text = h
        cell.paragraphs[0].runs[0].font.name = "Arial"
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(8.5)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=70, bottom=70, left=70, right=70)
    for i, item in enumerate(PRUEBAS_EVIDENCIAS_DATA):
        row = tbl_pru.rows[i+1]
        res_comp = f"Esp: {item['resultado_esperado'][:65]}...\nObt: {item['resultado_obtenido'][:65]}..."
        vals = [item["id"], item["nombre"], res_comp, item["estado"], item["evidencia"]]
        bg = "F8FAFC" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate(vals):
            cell = row.cells[j]
            cell.width = col_w_pru[j]
            cell.paragraphs[0].text = val
            cell.paragraphs[0].runs[0].font.name = "Arial"
            cell.paragraphs[0].runs[0].font.size = Pt(7.5)
            if j == 3:
                cell.paragraphs[0].runs[0].font.bold = True
                cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(22, 101, 52)
            elif j == 0:
                cell.paragraphs[0].runs[0].font.bold = True
                cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(30, 41, 59)
            else:
                cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(51, 65, 85)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=70, right=70)

    # --------------------------------------------------------------------------
    # 11. SEGURIDAD DE LA SOLUCIÓN
    # --------------------------------------------------------------------------
    add_h1("11. SEGURIDAD DE LA SOLUCIÓN INFORMÁTICA")
    add_h2("11.1 Matriz de Roles y Control de Acceso (RBAC)")
    tbl_rbac = doc.add_table(rows=len(SEGURIDAD_DATA["rbac_roles"])+1, cols=3)
    tbl_rbac.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_rbac, color="CBD5E1")
    col_w_rbac = [Inches(1.5), Inches(2.7), Inches(2.3)]
    headers_rbac = ["Rol Institucional", "Atribuciones Operativas Autorizadas", "Restricciones Estrictas de Seguridad"]
    for j, h in enumerate(headers_rbac):
        cell = tbl_rbac.rows[0].cells[j]
        cell.width = col_w_rbac[j]
        cell.paragraphs[0].text = h
        cell.paragraphs[0].runs[0].font.name = "Arial"
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(8.5)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
    for i, row_data in enumerate(SEGURIDAD_DATA["rbac_roles"]):
        row = tbl_rbac.rows[i+1]
        bg = "F8FAFC" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            cell.width = col_w_rbac[j]
            cell.paragraphs[0].text = val
            cell.paragraphs[0].runs[0].font.name = "Arial"
            cell.paragraphs[0].runs[0].font.size = Pt(8)
            cell.paragraphs[0].runs[0].font.bold = (j == 0)
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(51, 65, 85)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)

    add_h2("11.2 Mitigación de Vulnerabilidades Críticas (Top 10 OWASP)")
    tbl_owasp = doc.add_table(rows=len(SEGURIDAD_DATA["owasp_mitigaciones"])+1, cols=3)
    tbl_owasp.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_owasp, color="CBD5E1")
    col_w_owasp = [Inches(1.6), Inches(3.3), Inches(1.6)]
    headers_owasp = ["Vulnerabilidad OWASP", "Mecanismo de Mitigación en EduNova", "Capa de Aplicación"]
    for j, h in enumerate(headers_owasp):
        cell = tbl_owasp.rows[0].cells[j]
        cell.width = col_w_owasp[j]
        cell.paragraphs[0].text = h
        cell.paragraphs[0].runs[0].font.name = "Arial"
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(8.5)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
    for i, row_data in enumerate(SEGURIDAD_DATA["owasp_mitigaciones"]):
        row = tbl_owasp.rows[i+1]
        bg = "F8FAFC" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            cell.width = col_w_owasp[j]
            cell.paragraphs[0].text = val
            cell.paragraphs[0].runs[0].font.name = "Arial"
            cell.paragraphs[0].runs[0].font.size = Pt(8)
            cell.paragraphs[0].runs[0].font.bold = (j == 0)
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(51, 65, 85)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)

    # --------------------------------------------------------------------------
    # 12. PLAN DE OPERATIVIDAD Y MANTENIMIENTO
    # --------------------------------------------------------------------------
    add_h1("12. PLAN DE OPERATIVIDAD Y MANTENIMIENTO")
    add_h2("12.1 Mantenimiento Preventivo Periódico")
    for act in MANTENIMIENTO_DATA["preventivo"]:
        add_p(act, bold_prefix="• Tarea Preventiva: ")
    add_h2("12.2 Mantenimiento Correctivo y Contingencias")
    for act in MANTENIMIENTO_DATA["correctivo"]:
        add_p(act, bold_prefix="• Tarea Correctiva: ")
    add_h2("12.3 Procedimiento de Actualización Continua sin Caída de Servicio (Zero-Downtime)")
    for paso in MANTENIMIENTO_DATA["zero_downtime"]:
        add_p(paso)

    # --------------------------------------------------------------------------
    # 13. GESTIÓN DE RIESGOS TÉCNICOS Y OPERATIVOS
    # --------------------------------------------------------------------------
    add_h1("13. GESTIÓN DE RIESGOS TÉCNICOS Y OPERATIVOS")
    tbl_rsg = doc.add_table(rows=len(RIESGOS_DATA)+1, cols=5)
    tbl_rsg.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_rsg, color="CBD5E1")
    col_w_rsg = [Inches(0.6), Inches(1.8), Inches(1.0), Inches(1.6), Inches(1.5)]
    headers_rsg = ["ID", "Descripción del Riesgo", "Prob / Imp", "Medida Preventiva", "Plan de Contingencia"]
    for j, h in enumerate(headers_rsg):
        cell = tbl_rsg.rows[0].cells[j]
        cell.width = col_w_rsg[j]
        cell.paragraphs[0].text = h
        cell.paragraphs[0].runs[0].font.name = "Arial"
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(8.5)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
    for i, item in enumerate(RIESGOS_DATA):
        row = tbl_rsg.rows[i+1]
        vals = [item["id"], item["descripcion"], f"{item['probabilidad']} / {item['impacto']}", item["mitigacion"], item["contingencia"]]
        bg = "F8FAFC" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate(vals):
            cell = row.cells[j]
            cell.width = col_w_rsg[j]
            cell.paragraphs[0].text = val
            cell.paragraphs[0].runs[0].font.name = "Arial"
            cell.paragraphs[0].runs[0].font.size = Pt(7.5)
            cell.paragraphs[0].runs[0].font.bold = (j == 0)
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(51, 65, 85)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)

    # --------------------------------------------------------------------------
    # 14. CONCLUSIONES Y RECOMENDACIONES
    # --------------------------------------------------------------------------
    add_h1("14. CONCLUSIONES Y RECOMENDACIONES")
    add_h2("14.1 Conclusiones de Ingeniería y Gestión")
    for c in CONCLUSIONES_DATA:
        add_p(c, bold_prefix="• ")
    add_h2("14.2 Recomendaciones para la Operación Institucional")
    for r in RECOMENDACIONES_DATA:
        add_p(r, bold_prefix="• ")

    # --------------------------------------------------------------------------
    # 15. ACTA DE FIRMAS
    # --------------------------------------------------------------------------
    add_h1("15. ACTA DE CONTROL, RECEPCIÓN Y FIRMAS DE RESPONSABILIDAD")
    add_p("En señal de conformidad y validación técnica, se suscribe el presente informe administrativo:")
    tbl_fir = doc.add_table(rows=len(FIRMAS_DATA)+1, cols=3)
    tbl_fir.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_fir, color="CBD5E1")
    col_w_fir = [Inches(2.2), Inches(2.2), Inches(2.1)]
    headers_fir = ["Rol / Cargo Interviniente", "Institución / Entidad", "Firma y Responsabilidad Asumida"]
    for j, h in enumerate(headers_fir):
        cell = tbl_fir.rows[0].cells[j]
        cell.width = col_w_fir[j]
        cell.paragraphs[0].text = h
        cell.paragraphs[0].runs[0].font.name = "Arial"
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(8.5)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=80, bottom=80, left=90, right=90)
    for i, item in enumerate(FIRMAS_DATA):
        row = tbl_fir.rows[i+1]
        vals = [item["cargo"], item["entidad"], f"{item['responsabilidad']}\n\n___________________________\nFirma y Sello"]
        bg = "F8FAFC" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate(vals):
            cell = row.cells[j]
            cell.width = col_w_fir[j]
            cell.paragraphs[0].text = val
            cell.paragraphs[0].runs[0].font.name = "Arial"
            cell.paragraphs[0].runs[0].font.size = Pt(8)
            cell.paragraphs[0].runs[0].font.bold = (j == 0)
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(51, 65, 85)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)

    doc.save(ruta_salida)
    print(f"Documento DOCX generado exitosamente en: {ruta_salida}")

# ==============================================================================
# CONSTRUCTOR REPORTLAB (PDF)
# ==============================================================================
def generar_documento_pdf(ruta_salida):
    print("Iniciando construcción de archivo PDF con ReportLab...")
    doc = SimpleDocTemplate(
        ruta_salida,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )
    
    styles = getSampleStyleSheet()
    
    # Estilos personalizados
    c_sena = colors.HexColor("#39A900")
    c_slate = colors.HexColor("#1E293B")
    c_body = colors.HexColor("#334155")
    c_code = colors.HexColor("#E2E8F0")
    
    style_badge = ParagraphStyle(
        'CoverBadge',
        fontName='Arial-Bold',
        fontSize=10,
        leading=13,
        textColor=c_sena,
        spaceAfter=6
    )
    style_title = ParagraphStyle(
        'CoverTitle',
        fontName='Arial-Bold',
        fontSize=18,
        leading=22,
        textColor=c_slate,
        spaceAfter=8
    )
    style_sub = ParagraphStyle(
        'CoverSubtitle',
        fontName='Arial',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#475569"),
        spaceAfter=14
    )
    style_h1 = ParagraphStyle(
        'H1',
        fontName='Arial-Bold',
        fontSize=13,
        leading=16,
        textColor=c_sena,
        spaceBefore=12,
        spaceAfter=5,
        keepWithNext=True
    )
    style_h2 = ParagraphStyle(
        'H2',
        fontName='Arial-Bold',
        fontSize=10.5,
        leading=13,
        textColor=c_slate,
        spaceBefore=8,
        spaceAfter=3,
        keepWithNext=True
    )
    style_body = ParagraphStyle(
        'Body',
        fontName='Arial',
        fontSize=8.5,
        leading=11.5,
        textColor=c_body,
        spaceAfter=4
    )
    style_code = ParagraphStyle(
        'CodeBlock',
        fontName='Consolas',
        fontSize=7.5,
        leading=9.5,
        textColor=c_code
    )
    style_table_cell = ParagraphStyle(
        'TableCell',
        fontName='Arial',
        fontSize=7.5,
        leading=9.5,
        textColor=c_body
    )
    style_table_hdr = ParagraphStyle(
        'TableHdr',
        fontName='Arial-Bold',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.white
    )
    style_alert = ParagraphStyle(
        'AlertText',
        fontName='Arial',
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#166534")
    )
    
    story = []

    # PORTADA
    story.append(Spacer(1, 15))
    story.append(Paragraph("SERVICIO NACIONAL DE APRENDIZAJE – SENA | REGIONAL MAGDALENA", style_badge))
    story.append(Paragraph("INFORME ADMINISTRATIVO DE INSTALACIÓN, RESPALDO, MIGRACIÓN Y DESPLIEGUE A PRODUCCIÓN", style_title))
    story.append(Paragraph("SISTEMA DE INFORMACIÓN EDUNOVA — SISTEMA INTEGRADO DE EDUCACIÓN TÉCNICA (SINETEC)", style_sub))
    story.append(HRFlowable(width="100%", thickness=1.5, color=c_sena, spaceAfter=14))

    meta_rows = [
        [Paragraph("<b>Programa de Formación:</b>", style_table_cell), Paragraph(METADATOS["programa"], style_table_cell)],
        [Paragraph("<b>Código de Ficha / Proyecto:</b>", style_table_cell), Paragraph(f"Ficha {METADATOS['ficha']} | Proyecto SOFIA: {METADATOS['proyecto_sofia']}", style_table_cell)],
        [Paragraph("<b>Centro de Formación:</b>", style_table_cell), Paragraph(METADATOS["centro"], style_table_cell)],
        [Paragraph("<b>Nombre del Proyecto:</b>", style_table_cell), Paragraph(METADATOS["nombre_proyecto"], style_table_cell)],
        [Paragraph("<b>Versión del Documento:</b>", style_table_cell), Paragraph(METADATOS["version_documento"], style_table_cell)],
        [Paragraph("<b>Destinatario / Evaluación:</b>", style_table_cell), Paragraph(METADATOS["destinatario"], style_table_cell)],
        [Paragraph("<b>Autor / Desarrollador:</b>", style_table_cell), Paragraph(METADATOS["autor"], style_table_cell)],
        [Paragraph("<b>Ciudad y Fecha:</b>", style_table_cell), Paragraph(f"{METADATOS['ciudad']} — {METADATOS['fecha']}", style_table_cell)],
    ]
    t_meta = Table(meta_rows, colWidths=[150, 354])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(t_meta)
    story.append(PageBreak())

    # CONTROL DE VERSIONES
    story.append(Paragraph("CONTROL DE VERSIONES Y REGISTRO DOCUMENTAL", style_h1))
    story.append(Paragraph("El presente informe se rige por el estándar de trazabilidad y gestión de configuración documental del SENA:", style_body))
    ver_table_data = [[
        Paragraph(h, style_table_hdr) for h in ["Versión", "Fecha", "Autor", "Descripción del Cambio / Hito"]
    ]]
    for item in CONTROL_VERSIONES:
        ver_table_data.append([
            Paragraph(f"<b>{item['version']}</b>", style_table_cell),
            Paragraph(item["fecha"], style_table_cell),
            Paragraph(item["autor"], style_table_cell),
            Paragraph(item["descripcion"], style_table_cell)
        ])
    t_ver = Table(ver_table_data, colWidths=[55, 65, 120, 264])
    t_ver.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_slate),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
    ]))
    story.append(t_ver)
    story.append(Spacer(1, 10))

    # 1. INTRODUCCIÓN
    story.append(Paragraph("1. INTRODUCCIÓN Y CONTEXTO ADMINISTRATIVO", style_h1))
    for par in INTRODUCCION_TEXT.split("\n\n"):
        story.append(Paragraph(par, style_body))

    # Alert Box
    alert_data = [[Paragraph("<b>NOTA INSTITUCIONAL • DISTINCIÓN DE ESTADOS EN EL DOCUMENTO</b><br/>Conforme a lo instruido por la coordinación del SENA, este informe no simula servidores o dominios de producción ficticios. Se certifica que EduNova se encuentra desarrollado, estabilizado y verificado localmente con suites automatizadas (100% de éxito), y los procedimientos de producción se detallan como propuesta formal estandarizada pendiente de despliegue sobre el servidor final provisto por la institución educativa.", style_alert)]]
    t_alert = Table(alert_data, colWidths=[504])
    t_alert.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F0FDF4")),
        ('LINELEFT', (0, 0), (0, -1), 3, c_sena),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#BBF7D0")),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(t_alert)
    story.append(Spacer(1, 8))

    # 2. OBJETIVOS
    story.append(Paragraph("2. OBJETIVOS DEL INFORME", style_h1))
    story.append(Paragraph("<b>2.1 Objetivo General</b>", style_h2))
    story.append(Paragraph(OBJETIVOS_DATA["general"], style_body))
    story.append(Paragraph("<b>2.2 Objetivos Específicos</b>", style_h2))
    for obj in OBJETIVOS_DATA["especificos"]:
        story.append(Paragraph(f"• {obj}", style_body))

    # 3. ALCANCE
    story.append(Paragraph("3. ALCANCE DEL SISTEMA Y LÍMITES DEL DESPLIEGUE", style_h1))
    story.append(Paragraph("<b>3.1 Alcance Funcional de la Solución (Seis Paneles)</b>", style_h2))
    for p_fun in ALCANCE_DATA["funcional"]:
        parts = p_fun.split(": ")
        story.append(Paragraph(f"• <b>{parts[0]}:</b> {parts[1]}", style_body))
    story.append(Paragraph("<b>3.2 Alcance del Despliegue en Servidor</b>", style_h2))
    story.append(Paragraph(ALCANCE_DATA["despliegue"], style_body))
    story.append(Paragraph("<b>3.3 Límites y Exclusiones del Sistema</b>", style_h2))
    for lim in ALCANCE_DATA["limites"]:
        story.append(Paragraph(f"• <b>Exclusión:</b> {lim}", style_body))

    # 4. MATRIZ DE ESTADO
    story.append(PageBreak())
    story.append(Paragraph("4. MATRIZ DE ESTADO REAL DEL PROYECTO", style_h1))
    story.append(Paragraph("A continuación se consolida el estado verificado de cada componente técnico, distinguiendo entre lo realizado, propuesto y pendiente:", style_body))
    est_table_data = [[
        Paragraph(h, style_table_hdr) for h in ["Área Técnica", "Actividad / Componente", "Estado", "Evidencia"]
    ]]
    for item in MATRIZ_ESTADO:
        col = "#166534" if item["tipo"] == "REALIZADA" else ("#1E40AF" if item["tipo"] == "PROPUESTA" else "#B45309")
        est_table_data.append([
            Paragraph(item["area"], style_table_cell),
            Paragraph(item["actividad"], style_table_cell),
            Paragraph(f"<b><font color='{col}'>{item['tipo']}</font></b>", style_table_cell),
            Paragraph(item["evidencia"], style_table_cell)
        ])
    t_est = Table(est_table_data, colWidths=[100, 174, 80, 150])
    t_est.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_slate),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
    ]))
    story.append(t_est)
    story.append(Spacer(1, 10))

    # 5. REQUISITOS
    story.append(Paragraph("5. REQUISITOS DEL SERVIDOR (INFRAESTRUCTURA, HARDWARE, SOFTWARE Y RED)", style_h1))
    story.append(Paragraph("<b>5.1 Requisitos de Hardware del Servidor</b>", style_h2))
    hw_table_data = []
    for i, r in enumerate(REQUISITOS_SERVIDOR["hardware"]):
        st = style_table_hdr if i == 0 else style_table_cell
        hw_table_data.append([Paragraph(val, st) for val in r])
    t_hw = Table(hw_table_data, colWidths=[90, 115, 125, 174])
    t_hw.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_slate),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
    ]))
    story.append(t_hw)
    story.append(Spacer(1, 6))

    story.append(Paragraph("<b>5.2 Requisitos de Software Homologado</b>", style_h2))
    sw_table_data = []
    for i, r in enumerate(REQUISITOS_SERVIDOR["software"]):
        st = style_table_hdr if i == 0 else style_table_cell
        sw_table_data.append([Paragraph(val, st) for val in r])
    t_sw = Table(sw_table_data, colWidths=[95, 105, 110, 194])
    t_sw.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_slate),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
    ]))
    story.append(t_sw)
    story.append(Spacer(1, 6))

    story.append(Paragraph("<b>5.3 Requisitos de Red, Puertos y Firewall</b>", style_h2))
    net_table_data = []
    for i, r in enumerate(REQUISITOS_SERVIDOR["red_puertos"]):
        st = style_table_hdr if i == 0 else style_table_cell
        net_table_data.append([Paragraph(val, st) for val in r])
    t_net = Table(net_table_data, colWidths=[40, 50, 105, 105, 204])
    t_net.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_slate),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
    ]))
    story.append(t_net)

    # 6. PLAN DETALLADO DE INSTALACIÓN
    story.append(PageBreak())
    story.append(Paragraph("6. PLAN DETALLADO DE INSTALACIÓN Y DESPLIEGUE A PRODUCCIÓN", style_h1))
    story.append(Paragraph("Procedimiento paso a paso para la puesta en marcha técnica de EduNova sobre Ubuntu Server:", style_body))
    for fase in PLAN_INSTALACION_FASES:
        story.append(Paragraph(fase["fase"], style_h2))
        story.append(Paragraph(fase["descripcion"], style_body))
        
        # Bloque de código con fondo oscuro
        code_p = Paragraph("<br/>".join(fase["comandos"]).replace(" ", "&nbsp;"), style_code)
        t_code = Table([[code_p]], colWidths=[504])
        t_code.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#0F172A")),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ]))
        story.append(t_code)
        story.append(Spacer(1, 4))

    # 7. VARIABLES DE ENTORNO
    story.append(PageBreak())
    story.append(Paragraph("7. CONFIGURACIÓN DEL SISTEMA Y VARIABLES DE ENTORNO", style_h1))
    story.append(Paragraph("Variables de entorno gestionadas de forma segura mediante decouple en archivo .env:", style_body))
    env_table_data = [[
        Paragraph(h, style_table_hdr) for h in ["Variable", "Propósito en el Sistema", "Ejemplo", "Criticidad"]
    ]]
    for item in VARIABLES_ENTORNO_DATA:
        env_table_data.append([
            Paragraph(f"<b>{item[0]}</b>", style_table_cell),
            Paragraph(item[1], style_table_cell),
            Paragraph(item[2], style_table_cell),
            Paragraph(item[3], style_table_cell)
        ])
    t_env = Table(env_table_data, colWidths=[105, 174, 145, 80])
    t_env.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_slate),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
    ]))
    story.append(t_env)
    story.append(Spacer(1, 8))

    # 8. RESPALDO Y RECUPERACIÓN
    story.append(Paragraph("8. PLAN DE RESPALDO Y RECUPERACIÓN (BACKUP & DISASTER RECOVERY)", style_h1))
    story.append(Paragraph(PLAN_RESPALDO_DATA["descripcion"], style_body))
    story.append(Paragraph("<b>8.1 Política y Periodicidad de Respaldos</b>", style_h2))
    pol_table_data = []
    for i, r in enumerate(PLAN_RESPALDO_DATA["politica"]):
        st = style_table_hdr if i == 0 else style_table_cell
        pol_table_data.append([Paragraph(val, st) for val in r])
    t_pol = Table(pol_table_data, colWidths=[100, 95, 95, 95, 119])
    t_pol.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_slate),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
    ]))
    story.append(t_pol)
    story.append(Spacer(1, 6))

    story.append(Paragraph("<b>8.2 Script de Respaldo Automatizado (/usr/local/bin/backup_edunova.sh)</b>", style_h2))
    sh_p = Paragraph("<br/>".join(PLAN_RESPALDO_DATA["script_bash"][:18]).replace(" ", "&nbsp;"), style_code)
    t_sh = Table([[sh_p]], colWidths=[504])
    t_sh.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#0F172A")),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(t_sh)
    story.append(Paragraph(f"<b>Entrada Crontab (02:00 AM diario):</b> {PLAN_RESPALDO_DATA['cron_entry']}", style_body))

    story.append(Paragraph("<b>8.3 Protocolo de Restauración ante Desastres</b>", style_h2))
    for p in PLAN_RESPALDO_DATA["restauracion_pasos"]:
        story.append(Paragraph(p, style_body))

    story.append(Paragraph("<b>8.4 Inventario de Respaldos Reales en el Repositorio</b>", style_h2))
    bak_table_data = [[
        Paragraph(h, style_table_hdr) for h in ["Archivo Real", "Tamaño", "Ruta en Repositorio", "Descripción"]
    ]]
    for item in PLAN_RESPALDO_DATA["backups_reales"]:
        bak_table_data.append([
            Paragraph(f"<b>{item[0]}</b>", style_table_cell),
            Paragraph(item[1], style_table_cell),
            Paragraph(item[2], style_table_cell),
            Paragraph(item[3], style_table_cell)
        ])
    t_bak = Table(bak_table_data, colWidths=[130, 60, 140, 174])
    t_bak.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_slate),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
    ]))
    story.append(t_bak)

    # 9. MIGRACIÓN
    story.append(PageBreak())
    story.append(Paragraph("9. PLAN DE MIGRACIÓN DE DATOS", style_h1))
    story.append(Paragraph(PLAN_MIGRACION_DATA["metodologia"], style_body))
    story.append(Paragraph("<b>9.1 Etapas del Proceso de Migración (ETL)</b>", style_h2))
    for tit, desc in PLAN_MIGRACION_DATA["etapas"]:
        story.append(Paragraph(f"• <b>{tit}:</b> {desc}", style_body))
    story.append(Paragraph("<b>9.2 Chequeos y Verificaciones Post-Migración</b>", style_h2))
    for chk in PLAN_MIGRACION_DATA["chequeos_post_migracion"]:
        story.append(Paragraph(f"• <b>Verificación SQL:</b> {chk}", style_body))

    # 10. PRUEBAS
    story.append(Spacer(1, 6))
    story.append(Paragraph("10. PRUEBAS DE FUNCIONAMIENTO Y CONTROL DE CALIDAD", style_h1))
    story.append(Paragraph("Matriz de casos de prueba unitarias e integración ejecutadas y aprobadas (100% éxito):", style_body))
    pru_table_data = [[
        Paragraph(h, style_table_hdr) for h in ["ID", "Caso de Prueba", "Resultado Esperado vs Obtenido", "Estado", "Evidencia"]
    ]]
    for item in PRUEBAS_EVIDENCIAS_DATA:
        res_comp = f"<b>Esp:</b> {item['resultado_esperado'][:55]}...<br/><b>Obt:</b> {item['resultado_obtenido'][:55]}..."
        pru_table_data.append([
            Paragraph(f"<b>{item['id']}</b>", style_table_cell),
            Paragraph(item["nombre"], style_table_cell),
            Paragraph(res_comp, style_table_cell),
            Paragraph("<b><font color='#166534'>APROBADO</font></b>", style_table_cell),
            Paragraph(item["evidencia"], style_table_cell)
        ])
    t_pru = Table(pru_table_data, colWidths=[45, 115, 164, 60, 120])
    t_pru.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_slate),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 2.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
    ]))
    story.append(t_pru)

    # 11. SEGURIDAD
    story.append(PageBreak())
    story.append(Paragraph("11. SEGURIDAD DE LA SOLUCIÓN INFORMÁTICA", style_h1))
    story.append(Paragraph("<b>11.1 Matriz de Roles y Control de Acceso (RBAC)</b>", style_h2))
    rbac_table_data = [[
        Paragraph(h, style_table_hdr) for h in ["Rol Institucional", "Atribuciones Operativas", "Restricciones de Seguridad"]
    ]]
    for item in SEGURIDAD_DATA["rbac_roles"]:
        rbac_table_data.append([
            Paragraph(f"<b>{item[0]}</b>", style_table_cell),
            Paragraph(item[1], style_table_cell),
            Paragraph(item[2], style_table_cell)
        ])
    t_rbac = Table(rbac_table_data, colWidths=[110, 204, 190])
    t_rbac.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_slate),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
    ]))
    story.append(t_rbac)
    story.append(Spacer(1, 6))

    story.append(Paragraph("<b>11.2 Mitigación de Vulnerabilidades Críticas (Top 10 OWASP)</b>", style_h2))
    owasp_table_data = [[
        Paragraph(h, style_table_hdr) for h in ["Vulnerabilidad OWASP", "Mecanismo de Mitigación en EduNova", "Capa"]
    ]]
    for item in SEGURIDAD_DATA["owasp_mitigaciones"]:
        owasp_table_data.append([
            Paragraph(f"<b>{item[0]}</b>", style_table_cell),
            Paragraph(item[1], style_table_cell),
            Paragraph(item[2], style_table_cell)
        ])
    t_owasp = Table(owasp_table_data, colWidths=[120, 264, 120])
    t_owasp.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_slate),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 2.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
    ]))
    story.append(t_owasp)

    # 12. MANTENIMIENTO
    story.append(Spacer(1, 8))
    story.append(Paragraph("12. PLAN DE OPERATIVIDAD Y MANTENIMIENTO", style_h1))
    story.append(Paragraph("<b>12.1 Mantenimiento Preventivo Periódico</b>", style_h2))
    for act in MANTENIMIENTO_DATA["preventivo"]:
        story.append(Paragraph(f"• <b>Tarea Preventiva:</b> {act}", style_body))
    story.append(Paragraph("<b>12.2 Mantenimiento Correctivo y Contingencias</b>", style_h2))
    for act in MANTENIMIENTO_DATA["correctivo"]:
        story.append(Paragraph(f"• <b>Tarea Correctiva:</b> {act}", style_body))
    story.append(Paragraph("<b>12.3 Procedimiento de Actualización Continua (Zero-Downtime)</b>", style_h2))
    for paso in MANTENIMIENTO_DATA["zero_downtime"]:
        story.append(Paragraph(paso, style_body))

    # 13. GESTIÓN DE RIESGOS
    story.append(PageBreak())
    story.append(Paragraph("13. GESTIÓN DE RIESGOS TÉCNICOS Y OPERATIVOS", style_h1))
    rsg_table_data = [[
        Paragraph(h, style_table_hdr) for h in ["ID", "Descripción del Riesgo", "Prob / Imp", "Medida Preventiva", "Plan Contingencia"]
    ]]
    for item in RIESGOS_DATA:
        rsg_table_data.append([
            Paragraph(f"<b>{item['id']}</b>", style_table_cell),
            Paragraph(item["descripcion"], style_table_cell),
            Paragraph(f"{item['probabilidad']}<br/>{item['impacto']}", style_table_cell),
            Paragraph(item["mitigacion"], style_table_cell),
            Paragraph(item["contingencia"], style_table_cell)
        ])
    t_rsg = Table(rsg_table_data, colWidths=[35, 140, 65, 134, 130])
    t_rsg.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_slate),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
    ]))
    story.append(t_rsg)
    story.append(Spacer(1, 10))

    # 14. CONCLUSIONES Y RECOMENDACIONES
    story.append(Paragraph("14. CONCLUSIONES Y RECOMENDACIONES", style_h1))
    story.append(Paragraph("<b>14.1 Conclusiones de Ingeniería y Gestión</b>", style_h2))
    for c in CONCLUSIONES_DATA:
        story.append(Paragraph(f"• {c}", style_body))
    story.append(Paragraph("<b>14.2 Recomendaciones para la Operación Institucional</b>", style_h2))
    for r in RECOMENDACIONES_DATA:
        story.append(Paragraph(f"• {r}", style_body))

    # 15. ACTA DE FIRMAS
    story.append(Spacer(1, 8))
    story.append(Paragraph("15. ACTA DE CONTROL, RECEPCIÓN Y FIRMAS DE RESPONSABILIDAD", style_h1))
    story.append(Paragraph("En señal de conformidad y validación técnica, se suscribe el presente informe administrativo:", style_body))
    fir_table_data = [[
        Paragraph(h, style_table_hdr) for h in ["Rol / Cargo Interviniente", "Institución / Entidad", "Firma y Responsabilidad"]
    ]]
    for item in FIRMAS_DATA:
        fir_table_data.append([
            Paragraph(f"<b>{item['cargo']}</b>", style_table_cell),
            Paragraph(item["entidad"], style_table_cell),
            Paragraph(f"{item['responsabilidad']}<br/><br/>___________________________<br/><b>Firma y Sello</b>", style_table_cell)
        ])
    t_fir = Table(fir_table_data, colWidths=[150, 160, 194])
    t_fir.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_slate),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
    ]))
    story.append(t_fir)

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"Documento PDF generado exitosamente en: {ruta_salida}")

# ==============================================================================
# EJECUCIÓN PRINCIPAL
# ==============================================================================
if __name__ == "__main__":
    print("==================================================")
    print("GENERADOR MAESTRO DE INFORME ADMINISTRATIVO SENA")
    print("==================================================")
    generar_documento_docx(DOCX_PATH)
    generar_documento_pdf(PDF_PATH)
    print("==================================================")
    print("PROCESO DE GENERACIÓN FINALIZADO CON ÉXITO")
    print("==================================================")
