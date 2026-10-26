# -*- coding: utf-8 -*-
"""
Script Maestro de Generación del Manual de Usuario Oficial de EduNova
Genera:
- c:/EDUNOVA/EDUNOVA/docs/Manual_de_Usuario_EduNova.docx
- c:/EDUNOVA/EDUNOVA/docs/Manual_de_Usuario_EduNova.pdf
Cumple con los 20 criterios del Instrumento de Auditoría y Validación de Manuales de Usuario del SENA ADSI.
"""

import os
import sys
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS_DIR = os.path.join(BASE_DIR, 'docs')
os.makedirs(DOCS_DIR, exist_ok=True)

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, HRFlowable, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

try:
    pdfmetrics.registerFont(TTFont('Arial', 'C:/Windows/Fonts/arial.ttf'))
    pdfmetrics.registerFont(TTFont('Arial-Bold', 'C:/Windows/Fonts/arialbd.ttf'))
    pdfmetrics.registerFont(TTFont('Arial-Italic', 'C:/Windows/Fonts/ariali.ttf'))
except Exception:
    pass

# Helpers de formato Word
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_styled(doc, text, level=1):
    h = doc.add_paragraph()
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Arial"
    run.font.bold = True
    
    if level == 1:
        h.paragraph_format.space_before = Pt(20)
        h.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(139, 94, 60)  # Café #8B5E3C
        p_line = doc.add_paragraph()
        p_line.paragraph_format.space_before = Pt(0)
        p_line.paragraph_format.space_after = Pt(8)
        r_line = p_line.add_run("―" * 52)
        r_line.font.color.rgb = RGBColor(212, 163, 115)
    elif level == 2:
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(90, 56, 37)
    elif level == 3:
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(2)
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(45, 35, 30)
    return h

def add_p(doc, text, bold_prefix="", space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(45, 35, 30)
    r_body = p.add_run(text)
    r_body.font.name = "Arial"
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = RGBColor(55, 45, 40)
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(45, 35, 30)
    r_body = p.add_run(text)
    r_body.font.name = "Arial"
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = RGBColor(55, 45, 40)
    return p

def add_callout_box_docx(doc, title, text, box_type="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    bg_color = "F8F9FA"
    icon = "ℹ️ NOTA INFORMATIVA:"
    if box_type == "WARNING":
        bg_color = "FFFBEB"
        icon = "⚠️ ADVERTENCIA DE SEGURIDAD:"
    elif box_type == "IMPORTANT":
        bg_color = "FEF2F2"
        icon = "❗ REGLA INSTITUCIONAL OBLIGATORIA:"
    elif box_type == "TIP":
        bg_color = "F0FDF4"
        icon = "💡 CONSEJO PRÁCTICO:"
    
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=100, bottom=100, left=180, right=140)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(f"{icon} {title}\n")
    run_t.font.name = "Arial"
    run_t.font.size = Pt(9.5)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(45, 35, 30)
    
    run_b = p.add_run(text)
    run_b.font.name = "Arial"
    run_b.font.size = Pt(9)
    run_b.font.color.rgb = RGBColor(60, 50, 45)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)

def add_capture_marker_docx(doc, code, screen_name, profile, action, elements_to_highlight):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "EFF6FF")
    set_cell_margins(cell, top=110, bottom=110, left=180, right=140)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    run_tag = p.add_run(f"📷 MARCADOR DE CAPTURA REAL — [{code}: {screen_name}]\n")
    run_tag.font.name = "Arial"
    run_tag.font.size = Pt(9.5)
    run_tag.font.bold = True
    run_tag.font.color.rgb = RGBColor(29, 78, 216)
    
    details = (
        f"• Perfil a utilizar: {profile}\n"
        f"• Pantalla / URL: {screen_name}\n"
        f"• Acción a realizar para obtener la captura: {action}\n"
        f"• Instrucción visual (señalar con flechas o recuadros): {elements_to_highlight}"
    )
    run_d = p.add_run(details)
    run_d.font.name = "Arial"
    run_d.font.size = Pt(8.5)
    run_d.font.color.rgb = RGBColor(30, 58, 138)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)

def add_table_styled_docx(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    hdr_row = tbl.rows[0]
    for idx, heading in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_background(cell, "8B5E3C")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(heading)
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
    
    for r_idx, row_data in enumerate(data):
        row = tbl.rows[r_idx + 1]
        bg = "F8F9FA" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(str(val))
            r.font.name = "Arial"
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(45, 35, 30)
            
    if col_widths:
        for r in tbl.rows:
            for c_idx, w in enumerate(col_widths):
                r.cells[c_idx].width = Inches(w)
                
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)


# ==============================================================================
# REPORTLAB HELPERS
# ==============================================================================
class NumberedCanvasUser(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_decorations(self, page_count):
        if self._pageNumber == 1:
            return
        self.saveState()
        self.setFont("Arial", 8)
        self.setFillColor(colors.HexColor("#555555"))
        self.drawString(54, 11 * inch - 36, "EduNova | Manual Oficial de Usuario • SENA ADSI Ficha 3173430")
        self.setStrokeColor(colors.HexColor("#D4A373"))
        self.setLineWidth(0.6)
        self.line(54, 11 * inch - 40, 8.5 * inch - 54, 11 * inch - 40)
        
        page_text = f"Página {self._pageNumber} de {page_count}"
        self.drawRightString(8.5 * inch - 54, 36, page_text)
        self.drawString(54, 36, "Sistema de Información EduNova • Centro de Logística y Promoción Ecoturística")
        self.setStrokeColor(colors.HexColor("#E2E8F0"))
        self.setLineWidth(0.5)
        self.line(54, 46, 8.5 * inch - 54, 46)
        self.restoreState()


print("Compilando generador de manual de usuario...")
