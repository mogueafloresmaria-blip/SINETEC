from io import BytesIO

from openpyxl import Workbook
from django.test import TestCase
from django.core.files.uploadedfile import SimpleUploadedFile
from django.contrib.auth.models import User
from django.db.utils import IntegrityError
from instituciones.models import InstitucionEducativa
from academico.models import Competencia, Ficha, Matricula, ProgramaFormacion, ResultadoAprendizaje
from evaluaciones.models import JuicioEvaluativo
from seguimiento.models import BitacoraSeguimiento
from seguimiento.models import AsistenciaAprendiz
from usuarios.views import alertas_desercion_para_matricula
from .models import Rol, PerfilUsuario


class UsuarioModelTest(TestCase):
    def setUp(self):
        self.rol_coordinador = Rol.objects.create(
            nombre="Coordinador",
            descripcion="Líder del proceso de articulación con la Media Técnica"
        )
        self.user = User.objects.create_user(
            username="coordinador1",
            email="coord@sena.edu.co",
            first_name="Carlos",
            last_name="Mendoza"
        )
        self.perfil = PerfilUsuario.objects.get(usuario=self.user)
        self.perfil.rol = self.rol_coordinador
        self.perfil.numero_documento = "1082999888"
        self.perfil.save()

    def test_creacion_perfil_correcta(self):
        """Verifica que el perfil se haya creado y asociado correctamente al rol"""
        self.assertEqual(self.perfil.rol.nombre, "Coordinador")
        self.assertEqual(self.perfil.numero_documento, "1082999888")
        self.assertTrue(self.perfil.esta_activo)

    def test_estudiante_genera_carnet_qr_y_puede_escanearse(self):
        rol_estudiante = Rol.objects.create(nombre='Estudiante')
        aprendiz = User.objects.create_user(
            username='aprendiz_qr', first_name='Laura', last_name='QR'
        )
        perfil = PerfilUsuario.objects.get(usuario=aprendiz)
        perfil.rol = rol_estudiante
        perfil.numero_documento = '1002003012'
        perfil.save()

        self.assertIsNotNone(perfil.qr_token)
        self.assertTrue(perfil.qr_code.name)
        self.client.force_login(self.user)
        response = self.client.get(f'/estudiantes/qr/{perfil.qr_token}/')
        self.assertRedirects(response, f'/estudiantes/{perfil.pk}/')

    def test_lista_incluye_lector_qr_y_manifest_pwa(self):
        self.client.force_login(self.user)
        response = self.client.get('/estudiantes/')
        self.assertContains(response, 'Escanear carnet QR')
        self.assertContains(response, 'html5-qrcode')
        self.assertContains(response, 'manifest.webmanifest')

    def test_regla_rn001_unicidad_documento(self):
        """Verifica la regla RN-001: No pueden existir dos personas con el mismo número de documento"""
        otro_user = User.objects.create_user(
            username="instructor_dup",
            email="inst@sena.edu.co"
        )
        otro_perfil = PerfilUsuario.objects.get(usuario=otro_user)
        otro_perfil.numero_documento = "1082999888"  # Mismo número

        with self.assertRaises(IntegrityError):
            otro_perfil.save()

    def test_lista_importa_excel_y_asigna_ficha_colegio(self):
        institucion = InstitucionEducativa.objects.create(
            codigo_dane='147001009999', nombre='Colegio Excel', municipio='Santa Marta'
        )
        programa = ProgramaFormacion.objects.create(
            codigo_programa='EXCEL-01', denominacion='Programa Excel'
        )
        ficha = Ficha.objects.create(
            codigo_ficha='EXCEL-2501', programa=programa, institucion=institucion,
            instructor_lider=self.user, fecha_inicio='2026-02-01', fecha_fin='2026-11-30',
        )
        libro = Workbook()
        hoja = libro.active
        hoja.append(['tipo_documento', 'numero_documento', 'nombres', 'apellidos', 'correo', 'grado_escolar', 'codigo_ficha'])
        hoja.append(['TI', '1002003010', 'Maria', 'Excel', 'maria.excel@example.com', '10', ficha.codigo_ficha])
        contenido = BytesIO()
        libro.save(contenido)
        archivo = SimpleUploadedFile('estudiantes.xlsx', contenido.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        self.client.force_login(self.user)
        response = self.client.post('/estudiantes/', {'accion': 'importar_excel', 'archivo': archivo})

        self.assertRedirects(response, '/estudiantes/')
        perfil = PerfilUsuario.objects.get(numero_documento='1002003010')
        matricula = Matricula.objects.get(aprendiz=perfil.usuario)
        self.assertEqual(matricula.ficha, ficha)
        self.assertEqual(matricula.ficha.institucion, institucion)

    def test_lista_rechaza_excel_con_ficha_inexistente_sin_crear(self):
        libro = Workbook()
        hoja = libro.active
        hoja.append(['tipo_documento', 'numero_documento', 'nombres', 'apellidos', 'correo', 'grado_escolar', 'codigo_ficha'])
        hoja.append(['TI', '1002003011', 'Aprendiz', 'Invalido', 'invalido@example.com', '10', 'NO-EXISTE'])
        contenido = BytesIO()
        libro.save(contenido)
        archivo = SimpleUploadedFile('invalido.xlsx', contenido.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        self.client.force_login(self.user)
        response = self.client.post('/estudiantes/', {'accion': 'importar_excel', 'archivo': archivo})

        self.assertEqual(response.status_code, 200)
        self.assertFalse(PerfilUsuario.objects.filter(numero_documento='1002003011').exists())
        self.assertContains(response, 'no existe la ficha')

    def test_ficha_estudiante_muestra_historial_y_acudiente(self):
        rol_estudiante = Rol.objects.create(nombre='Estudiante')
        estudiante_user = User.objects.create_user(
            username='aprendiz_ficha',
            email='aprendiz@example.com',
            first_name='Ana',
            last_name='Torres',
        )
        estudiante = PerfilUsuario.objects.get(usuario=estudiante_user)
        estudiante.rol = rol_estudiante
        estudiante.numero_documento = '1002003007'
        estudiante.save()

        institucion = InstitucionEducativa.objects.create(
            codigo_dane='147001000888',
            nombre='Institución de Prueba',
            municipio='Santa Marta',
        )
        programa = ProgramaFormacion.objects.create(
            codigo_programa='228199',
            denominacion='Técnico en Sistemas',
        )
        ficha = Ficha.objects.create(
            codigo_ficha='2501999',
            programa=programa,
            institucion=institucion,
            instructor_lider=self.user,
            fecha_inicio='2026-02-01',
            fecha_fin='2026-11-30',
        )
        matricula = Matricula.objects.create(
            ficha=ficha,
            aprendiz=estudiante_user,
            grado_escolar='10',
            acudiente_nombre='Carlos Torres',
            acudiente_telefono='3001234567',
        )
        competencia = Competencia.objects.create(
            programa=programa,
            codigo='220501001',
            descripcion='Competencia de prueba',
        )
        rap = ResultadoAprendizaje.objects.create(
            competencia=competencia,
            codigo='RAP-01',
            descripcion='Resultado de prueba',
        )
        JuicioEvaluativo.objects.create(
            matricula=matricula,
            resultado_aprendizaje=rap,
            instructor=self.user,
            juicio_valor='A',
            fecha_evaluacion='2026-09-15',
        )
        BitacoraSeguimiento.objects.create(
            ficha=ficha,
            matricula=matricula,
            instructor=self.user,
            fecha_visita='2026-09-15',
            tipo_seguimiento='Presencial Aula',
            observaciones='Asistencia registrada en la sesión.',
        )

        self.client.force_login(self.user)
        response = self.client.get(f'/estudiantes/{estudiante.pk}/')

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Carlos Torres')
        self.assertContains(response, 'RAP-01')
        self.assertContains(response, 'Asistencia registrada en la sesión.')

        pdf_response = self.client.get(f'/estudiantes/{estudiante.pk}/boletin-pdf/')
        self.assertEqual(pdf_response.status_code, 200)
        self.assertEqual(pdf_response['Content-Type'], 'application/pdf')
        self.assertTrue(pdf_response.content.startswith(b'%PDF-1.4'))
        self.assertIn(b'SINETEC - BOLETIN OFICIAL', pdf_response.content)

    def test_alerta_por_cuatro_ausencias_consecutivas(self):
        rol_estudiante = Rol.objects.create(nombre='Estudiante')
        estudiante_user = User.objects.create_user(username='aprendiz_alerta_fallas')
        perfil = PerfilUsuario.objects.get(usuario=estudiante_user)
        perfil.rol = rol_estudiante
        perfil.numero_documento = '1002003008'
        perfil.save()
        institucion = InstitucionEducativa.objects.create(
            codigo_dane='147001000889', nombre='Colegio Alertas', municipio='Santa Marta'
        )
        programa = ProgramaFormacion.objects.create(codigo_programa='228198', denominacion='Programa Alertas')
        ficha = Ficha.objects.create(
            codigo_ficha='2501998', programa=programa, institucion=institucion,
            instructor_lider=self.user, fecha_inicio='2026-02-01', fecha_fin='2026-11-30',
        )
        matricula = Matricula.objects.create(ficha=ficha, aprendiz=estudiante_user, grado_escolar='10')
        for dia in range(11, 15):
            AsistenciaAprendiz.objects.create(
                matricula=matricula, fecha=f'2026-09-{dia}', estado='A', registrado_por=self.user
            )

        alertas = alertas_desercion_para_matricula(matricula)

        self.assertTrue(any(alerta['tipo'] == 'inasistencia' for alerta in alertas))

    def test_alerta_por_dos_rap_no_aprobados(self):
        rol_estudiante = Rol.objects.create(nombre='Estudiante')
        estudiante_user = User.objects.create_user(username='aprendiz_alerta_rap')
        perfil = PerfilUsuario.objects.get(usuario=estudiante_user)
        perfil.rol = rol_estudiante
        perfil.numero_documento = '1002003009'
        perfil.save()
        institucion = InstitucionEducativa.objects.create(
            codigo_dane='147001000890', nombre='Colegio RAP', municipio='Santa Marta'
        )
        programa = ProgramaFormacion.objects.create(codigo_programa='228197', denominacion='Programa RAP')
        ficha = Ficha.objects.create(
            codigo_ficha='2501997', programa=programa, institucion=institucion,
            instructor_lider=self.user, fecha_inicio='2026-02-01', fecha_fin='2026-11-30',
        )
        matricula = Matricula.objects.create(ficha=ficha, aprendiz=estudiante_user, grado_escolar='10')
        competencia = Competencia.objects.create(programa=programa, codigo='COMP-01', descripcion='Competencia')
        for codigo in ('RAP-01', 'RAP-02'):
            rap = ResultadoAprendizaje.objects.create(competencia=competencia, codigo=codigo, descripcion='Resultado')
            JuicioEvaluativo.objects.create(
                matricula=matricula, resultado_aprendizaje=rap, instructor=self.user,
                juicio_valor='D', fecha_evaluacion='2026-09-15',
            )

        alertas = alertas_desercion_para_matricula(matricula)

        self.assertTrue(any(alerta['tipo'] == 'rendimiento' for alerta in alertas))

    def test_modulo_alertas_tempranas_renderiza(self):
        self.client.force_login(self.user)

        response = self.client.get('/alertas/')

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Alertas Tempranas y Deserción')
        self.assertContains(response, "dos o más juicios 'D'")

    def test_descargar_plantilla_instructores(self):
        self.client.force_login(self.user)
        response = self.client.get('/usuarios/plantilla-instructores/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('plantilla_registro_masivo_instructores_sena.xlsx', response['Content-Disposition'])

    def test_importar_instructores_masivo_excel(self):
        self.client.force_login(self.user)
        wb = Workbook()
        ws = wb.active
        ws.append(["tipo_documento", "numero_documento", "nombres", "apellidos", "correo", "telefono", "especialidad"])
        ws.append(["CC", "1098765432", "Marcos", "Pacheco", "mpacheco@sena.edu.co", "3015554433", "Sistemas"])

        archivo_bytes = BytesIO()
        wb.save(archivo_bytes)
        archivo_bytes.seek(0)
        archivo = SimpleUploadedFile(
            "instructores_test.xlsx",
            archivo_bytes.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        response = self.client.post('/usuarios/importar-instructores/', {'archivo': archivo}, follow=True)
        self.assertEqual(response.status_code, 200)

        # Verificar que el usuario e instructor existen
        user_inst = User.objects.get(username="inst_1098765432")
        self.assertEqual(user_inst.first_name, "Marcos")
        self.assertEqual(user_inst.last_name, "Pacheco")
        self.assertEqual(user_inst.email, "mpacheco@sena.edu.co")
        self.assertTrue(user_inst.check_password("Sena5432*"))

        perfil = user_inst.perfil
        self.assertEqual(perfil.numero_documento, "1098765432")
        self.assertEqual(perfil.tipo_documento, "CC")
        self.assertIn("Instructor", perfil.rol.nombre)

    def test_importar_instructores_masivo_duplicado_rn001(self):
        self.client.force_login(self.user)
        # Intentar importar el mismo documento de self.user ("1082999888")
        wb = Workbook()
        ws = wb.active
        ws.append(["tipo_documento", "numero_documento", "nombres", "apellidos", "correo", "telefono", "especialidad"])
        ws.append(["CC", "1082999888", "Carlos Duplicado", "Mendoza", "cdup@sena.edu.co", "3000000000", "ADSO"])

        archivo_bytes = BytesIO()
        wb.save(archivo_bytes)
        archivo_bytes.seek(0)
        archivo = SimpleUploadedFile(
            "instructores_dup.xlsx",
            archivo_bytes.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        response = self.client.post('/usuarios/importar-instructores/', {'archivo': archivo}, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "RN-001")
        self.assertFalse(User.objects.filter(username="inst_1082999888").exists())

    def test_aprendiz_restringido_de_coordinacion_y_usuarios(self):
        rol_aprendiz, _ = Rol.objects.get_or_create(nombre='Aprendiz')
        user_aprendiz = User.objects.create_user(
            username='aprendiz_seguro',
            password='Password123*',
            first_name='Pedro',
            last_name='Segura'
        )
        perfil_ap = user_aprendiz.perfil
        perfil_ap.rol = rol_aprendiz
        perfil_ap.numero_documento = '1033445566'
        perfil_ap.save()

        self.client.force_login(user_aprendiz)
        # Intentar acceder al panel de coordinación
        res_coord = self.client.get('/coordinacion/')
        self.assertRedirects(res_coord, '/aprendiz/')

        # Intentar acceder a la lista de estudiantes
        res_est = self.client.get('/estudiantes/')
        self.assertRedirects(res_est, '/aprendiz/')

    def test_qr_antifraude_rotacion_diaria(self):
        from django.utils import timezone
        hoy = timezone.localdate()
        rol_estudiante, _ = Rol.objects.get_or_create(nombre='Estudiante')
        user_ap = User.objects.create_user(
            username='aprendiz_antifraude',
            first_name='Ana',
            last_name='Fraude'
        )
        perfil_ap = user_ap.perfil
        perfil_ap.rol = rol_estudiante
        perfil_ap.numero_documento = '1099887766'
        perfil_ap.save()

        # 1. Escaneo con fecha anterior (simula foto de ayer tomada en casa)
        ayer = hoy - timezone.timedelta(days=1)
        res_ayer = self.client.get(f'/estudiantes/qr/{perfil_ap.qr_token}/?dia={ayer}')
        self.assertEqual(res_ayer.status_code, 410)
        self.assertContains(res_ayer, 'Código QR vencido o fotografía estática no admitida', status_code=410)

        # 2. Escaneo con fecha de hoy
        res_hoy = self.client.get(f'/estudiantes/qr/{perfil_ap.qr_token}/?dia={hoy}')
        self.assertEqual(res_hoy.status_code, 200)

    def test_convocatorias_incluyen_alimentacion_transporte_contrato(self):
        response = self.client.get('/')
        self.assertEqual(response.status_code, 200)
        self.assertIn('convocatorias', response.context)
        convs = response.context['convocatorias']
        keys = [c['id'] for c in convs]
        self.assertIn('alimentacion', keys)
        self.assertIn('transporte', keys)
        self.assertIn('contrato', keys)

