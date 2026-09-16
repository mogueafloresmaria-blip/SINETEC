from io import BytesIO

from openpyxl import load_workbook
from django.test import TestCase, Client
from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from instituciones.models import InstitucionEducativa
from academico.models import ProgramaFormacion, Competencia, ResultadoAprendizaje, Ficha, Matricula
from .models import JuicioEvaluativo


class EvaluacionesModelTest(TestCase):
    def setUp(self):
        self.colegio = InstitucionEducativa.objects.create(
            codigo_dane="147001000777",
            nombre="I.E. Agropecuaria",
            municipio="Aracataca"
        )
        self.programa = ProgramaFormacion.objects.create(
            codigo_programa="228106",
            denominacion="Técnico en Sistemas"
        )
        self.competencia = Competencia.objects.create(
            programa=self.programa,
            codigo="220501001",
            descripcion="Mantenimiento Preventivo y Correctivo de Equipos"
        )
        self.rap = ResultadoAprendizaje.objects.create(
            competencia=self.competencia,
            codigo="RAP-01",
            descripcion="Desensamblar y ensamblar equipos de cómputo de acuerdo con el manual."
        )
        self.instructor = User.objects.create_user(
            username="inst_roberto",
            first_name="Roberto",
            last_name="Peña"
        )
        self.ficha = Ficha.objects.create(
            codigo_ficha="2501452",
            programa=self.programa,
            institucion=self.colegio,
            instructor_lider=self.instructor,
            fecha_inicio="2026-02-01",
            fecha_fin="2027-11-30",
            periodo_cerrado=False
        )
        self.aprendiz = User.objects.create_user(
            username="aprendiz_camila",
            first_name="Camila",
            last_name="Ortiz"
        )
        self.matricula = Matricula.objects.create(
            ficha=self.ficha,
            aprendiz=self.aprendiz,
            grado_escolar="10"
        )

    def test_regla_rn003_juicio_aprobado_valido(self):
        """Verifica la asignación de un juicio válido ('A')"""
        juicio = JuicioEvaluativo(
            matricula=self.matricula,
            resultado_aprendizaje=self.rap,
            instructor=self.instructor,
            juicio_valor='A',
            fecha_evaluacion="2026-09-14"
        )
        juicio.full_clean()
        juicio.save()
        self.assertEqual(juicio.juicio_valor, 'A')

    def test_regla_rn003_rechazar_juicio_invalido(self):
        """Verifica la regla RN-003: Rechaza juicios no válidos distintos de 'A' o 'D'"""
        juicio = JuicioEvaluativo(
            matricula=self.matricula,
            resultado_aprendizaje=self.rap,
            instructor=self.instructor,
            juicio_valor='B',  # Valor no permitido
            fecha_evaluacion="2026-09-14"
        )
        with self.assertRaises(ValidationError):
            juicio.full_clean()

    def test_regla_rn004_bloqueo_por_periodo_cerrado(self):
        """Verifica la regla RN-004: Impide evaluar si la ficha tiene su periodo formalmente cerrado"""
        self.ficha.periodo_cerrado = True
        self.ficha.save()

        juicio = JuicioEvaluativo(
            matricula=self.matricula,
            resultado_aprendizaje=self.rap,
            instructor=self.instructor,
            juicio_valor='A',
            fecha_evaluacion="2026-09-14"
        )
        with self.assertRaises(ValidationError):
            juicio.full_clean()

    def test_endpoint_raps_por_ficha_filtra_por_programa(self):
        otro_programa = ProgramaFormacion.objects.create(
            codigo_programa="228107",
            denominacion="Técnico en Administración"
        )
        otra_competencia = Competencia.objects.create(
            programa=otro_programa,
            codigo="210601001",
            descripcion="Gestionar procesos administrativos"
        )
        otro_rap = ResultadoAprendizaje.objects.create(
            competencia=otra_competencia,
            codigo="RAP-OTRO",
            descripcion="Resultado de otro programa"
        )

        client = Client()
        client.force_login(self.instructor)
        response = client.get('/evaluaciones/raps/', {'ficha_id': self.ficha.id})

        self.assertEqual(response.status_code, 200)
        self.assertEqual([rap['id'] for rap in response.json()['raps']], [self.rap.id])
        self.assertNotIn(otro_rap.id, [rap['id'] for rap in response.json()['raps']])

    def test_endpoint_aprendices_por_rap_devuelve_lista_y_juicio(self):
        JuicioEvaluativo.objects.create(
            matricula=self.matricula,
            resultado_aprendizaje=self.rap,
            instructor=self.instructor,
            juicio_valor='A',
            fecha_evaluacion='2026-09-14',
        )
        client = Client()
        client.force_login(self.instructor)

        response = client.get('/evaluaciones/aprendices/', {
            'ficha_id': self.ficha.id,
            'rap_id': self.rap.id,
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['aprendices'][0]['matricula_id'], self.matricula.id)
        self.assertEqual(response.json()['aprendices'][0]['juicio'], 'A')

    def test_exportar_sabana_excel_aplica_periodo_y_trimestre(self):
        juicio = JuicioEvaluativo.objects.create(
            matricula=self.matricula,
            resultado_aprendizaje=self.rap,
            instructor=self.instructor,
            juicio_valor='A',
            fecha_evaluacion='2026-09-14',
        )
        client = Client()
        client.force_login(self.instructor)
        response = client.get('/evaluaciones/exportar-excel/', {
            'ficha': self.ficha.id,
            'rap': self.rap.id,
            'periodo': '2026',
            'trimestre': '3',
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        libro = load_workbook(BytesIO(response.content))
        filas = list(libro.active.iter_rows(values_only=True))
        self.assertEqual(filas[1][0], self.ficha.codigo_ficha)
        self.assertEqual(filas[1][5], 'A - Aprobado (Alcanzó los criterios de desempeño)')
        self.assertEqual(filas[1][7], '2026-09-14')
